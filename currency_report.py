"""
currency_report.py - Standalone Currency Rate report.

Emits a dedicated xlsx with:
    Current Date block  | Previous Date block | Zenith Value
    Currency | BDT rate | Currency | BDT rate | Zenith

Rules:
- Sorted descending by current BDT rate.
- Current-day cells highlighted (cyan fill) only where value differs from previous.
- For USD, a "(DD-MMM-YY)" suffix is rendered next to the current value, where the
  date is the most recent date USD ever changed (tracked in usd_tracker.json).
  The suffix text is colored red; the USD value itself is unhighlighted unless
  the rate changed from previous day.
- Zenith Value = 1 / current_rate.
"""

from __future__ import annotations

import logging
import os
from dataclasses import dataclass
from datetime import date, datetime
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from currency_archive import (
    RateSnapshot,
    UsdTracker,
    latest_prior_snapshot,
    load_usd_tracker,
    save_snapshot,
    update_usd_tracker,
)

logger = logging.getLogger("travelport.currency_report")

# ── Styles ──────────────────────────────────────────────────────────────────
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
HEADER_FONT = Font(name="Calibri", bold=True, size=13)
DATE_FONT = Font(name="Calibri", bold=True, size=12)
TABLE_HEADER_FONT = Font(name="Calibri", bold=True, size=12)
BODY_FONT = Font(name="Calibri", size=11)
ZENITH_HEADER_FONT = Font(name="Calibri", bold=True, size=12, color="B8860B")
USD_DATE_FONT_COLOR = "B80000"  # red used only for the USD bracket

CHANGED_FILL = PatternFill(start_color="C6E7F5", end_color="C6E7F5", fill_type="solid")
CENTER = Alignment(horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")


@dataclass(frozen=True)
class CurrencyReportResult:
    path: str
    run_date: date
    previous_date: Optional[date]
    usd_tracker: UsdTracker


# ── Public API ──────────────────────────────────────────────────────────────


def generate_currency_report(
    current_rates: dict,
    output_path: str,
    run_date: Optional[date] = None,
    persist: bool = True,
) -> CurrencyReportResult:
    """Build the Currency Rate xlsx.

    current_rates : {"USD": 122.94, "KWD": 398.42704, ...}  (BDT per 1 unit)
    output_path   : absolute path for the xlsx (overwritten on save).
    run_date      : effective date for the "Current Date" column; defaults to today.
    persist       : when True, saves today's snapshot and updates the USD tracker.
    """
    run_date = run_date or datetime.now().date()
    normalized = {str(k).upper(): float(v) for k, v in current_rates.items() if v}

    if persist:
        save_snapshot(RateSnapshot(snapshot_date=run_date, rates=normalized, source="live"))

    prior = latest_prior_snapshot(run_date)
    prev_rates = prior.rates if prior else {}

    usd_rate = normalized.get("USD")
    if persist and usd_rate:
        tracker = update_usd_tracker(usd_rate, run_date)
    else:
        tracker = load_usd_tracker() or UsdTracker(
            rate=usd_rate or 0.0, last_changed_date=run_date
        )

    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Currency Rate")
    _write_sheet(
        ws,
        current_rates=normalized,
        previous_rates=prev_rates,
        run_date=run_date,
        previous_date=prior.snapshot_date if prior else None,
        usd_tracker=tracker,
    )

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)
    return CurrencyReportResult(
        path=output_path,
        run_date=run_date,
        previous_date=prior.snapshot_date if prior else None,
        usd_tracker=tracker,
    )


# ── Worksheet layout ────────────────────────────────────────────────────────


def _write_sheet(
    ws,
    current_rates: dict,
    previous_rates: dict,
    run_date: date,
    previous_date: Optional[date],
    usd_tracker: UsdTracker,
) -> None:
    # Three separate blocks with narrow spacer columns between them.
    #   A: Currency  | B: Exchange Rate To BDT    (Current block)
    #   C: spacer
    #   D: Currency  | E: Exchange Rate To BDT    (Previous block)
    #   F: spacer
    #   G: Currency  | H: Value                   (Zenith block)
    currencies = sorted(
        current_rates.keys(), key=lambda c: current_rates[c], reverse=True
    )

    # Row 1 — block titles, each merged across its 2 columns.
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    cur_title = ws.cell(row=1, column=1, value="Current Date")
    cur_title.font = HEADER_FONT
    cur_title.alignment = CENTER

    ws.merge_cells(start_row=1, start_column=4, end_row=1, end_column=5)
    prev_title = ws.cell(row=1, column=4, value="Previous Date")
    prev_title.font = HEADER_FONT
    prev_title.alignment = CENTER

    # Zenith title spans rows 1–2 so its sub-header aligns with the others on row 3.
    ws.merge_cells(start_row=1, start_column=7, end_row=2, end_column=8)
    zen_title = ws.cell(row=1, column=7, value="Zenith")
    zen_title.font = ZENITH_HEADER_FONT
    zen_title.alignment = CENTER

    # Row 2 — date strings under Current/Previous (merged across their 2 cols).
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2)
    cur_date_cell = ws.cell(row=2, column=1, value=run_date.strftime("%d-%b-%y"))
    cur_date_cell.font = DATE_FONT
    cur_date_cell.alignment = CENTER

    prev_label_txt = previous_date.strftime("%d-%b-%y") if previous_date else "—"
    ws.merge_cells(start_row=2, start_column=4, end_row=2, end_column=5)
    prev_date_cell = ws.cell(row=2, column=4, value=prev_label_txt)
    prev_date_cell.font = DATE_FONT
    prev_date_cell.alignment = CENTER

    # Row 3 — sub-headers for every block.
    subheaders = [
        (1, "Currency"),
        (2, "Exchange Rate To BDT"),
        (4, "Currency"),
        (5, "Exchange Rate To BDT"),
        (7, "Currency"),
        (8, "Value"),
    ]
    for col, text in subheaders:
        cell = ws.cell(row=3, column=col, value=text)
        cell.font = TABLE_HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER

    # Borders for the merged header/date cells (openpyxl only styles the anchor cell).
    header_ranges = [
        (1, 1, 1, 2),  # Current title
        (1, 4, 1, 5),  # Previous title
        (1, 7, 2, 8),  # Zenith title
        (2, 1, 2, 2),  # Current date
        (2, 4, 2, 5),  # Previous date
    ]
    for r1, c1, r2, c2 in header_ranges:
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                ws.cell(row=r, column=c).border = THIN_BORDER

    # Data rows
    row = 4
    for cur in currencies:
        cur_rate = current_rates[cur]
        prev_rate = previous_rates.get(cur)
        is_changed = prev_rate is None or not _rates_equal(cur_rate, prev_rate)

        # Current block — highlight on change.
        cur_label = ws.cell(row=row, column=1, value=f"1 {cur}")
        cur_label.font = BODY_FONT
        cur_label.alignment = CENTER
        cur_label.border = THIN_BORDER

        cur_value = ws.cell(row=row, column=2)
        cur_value.border = THIN_BORDER
        cur_value.alignment = CENTER

        if is_changed:
            cur_label.fill = CHANGED_FILL
            cur_value.fill = CHANGED_FILL

        if cur == "USD":
            _write_usd_current_cell(cur_value, cur_rate, usd_tracker, run_date)
        else:
            cur_value.value = cur_rate
            cur_value.font = BODY_FONT
            cur_value.number_format = "0.000000"

        # Previous block — no highlight.
        prev_label_cell = ws.cell(row=row, column=4, value=f"1 {cur}")
        prev_label_cell.font = BODY_FONT
        prev_label_cell.alignment = CENTER
        prev_label_cell.border = THIN_BORDER

        prev_value_cell = ws.cell(row=row, column=5)
        prev_value_cell.border = THIN_BORDER
        prev_value_cell.alignment = CENTER
        prev_value_cell.font = BODY_FONT
        if prev_rate is not None:
            prev_value_cell.value = prev_rate
            prev_value_cell.number_format = "0.000000"
        else:
            prev_value_cell.value = "—"

        # Zenith block — 1 / current_rate, no highlight.
        zen_label = ws.cell(row=row, column=7, value=f"1 {cur}")
        zen_label.font = BODY_FONT
        zen_label.alignment = CENTER
        zen_label.border = THIN_BORDER

        zen_value = ws.cell(row=row, column=8, value=1.0 / cur_rate if cur_rate else 0)
        zen_value.font = BODY_FONT
        zen_value.alignment = RIGHT
        zen_value.border = THIN_BORDER
        zen_value.number_format = "0.000000"

        row += 1

    # Column widths — two data columns per block + narrow spacers.
    widths = {
        "A": 10, "B": 24,
        "C": 2,
        "D": 10, "E": 24,
        "F": 2,
        "G": 10, "H": 14,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _write_usd_current_cell(cell, rate: float, tracker: UsdTracker, run_date: date) -> None:
    """Render USD current-day value with `(DD-MMM-YY)` suffix.

    Uses rich-text (mixed colors) when openpyxl's lxml writer is available, so the
    rate stays black/bold and only the date bracket is red. Falls back to a plain
    string with the whole cell red when lxml is missing — openpyxl's pure-python
    CellRichText writer has a known bug that corrupts the xlsx on save.
    """
    date_str = tracker.last_changed_date.strftime("%d-%b-%y").upper()
    rate_str = _format_rate(rate)

    if _has_lxml_rich_text():
        try:
            from openpyxl.cell.rich_text import CellRichText, TextBlock
            from openpyxl.cell.text import InlineFont

            cell.value = CellRichText(
                TextBlock(InlineFont(rFont="Calibri", sz=11, b=True), f"{rate_str}  "),
                TextBlock(
                    InlineFont(rFont="Calibri", sz=11, b=True, color=USD_DATE_FONT_COLOR),
                    f"({date_str})",
                ),
            )
            return
        except Exception as exc:  # pragma: no cover - defensive
            logger.debug("  Rich-text USD cell failed (%s); using plain fallback.", exc)

    cell.value = f"{rate_str}  ({date_str})"
    cell.font = Font(name="Calibri", size=11, bold=True, color=USD_DATE_FONT_COLOR)


def _has_lxml_rich_text() -> bool:
    """True when lxml is importable — required for openpyxl rich-text writes."""
    try:
        import lxml  # noqa: F401
    except ImportError:
        return False
    return True


def _format_rate(rate: float) -> str:
    # Trim trailing zeros but keep at least 2 decimals, match image style.
    formatted = f"{rate:.6f}".rstrip("0").rstrip(".")
    if "." not in formatted:
        formatted += ".00"
    elif len(formatted.split(".")[1]) < 2:
        formatted = f"{rate:.2f}"
    return formatted


def _rates_equal(a: float, b: float, tol: float = 1e-6) -> bool:
    return abs(a - b) <= tol
