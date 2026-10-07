"""gui_support.py - Pure helpers behind the GUI (no Tk), so they can be tested.

Route/airline choices for the pickers, filter checks as the user types, and
reading a fare report's Re-run Needed list.
"""

from __future__ import annotations

import glob
import os

from excel_report import RERUN_SHEET, _airline_sort_key
from exceptions import ValidationError
from validators import validate_airline_code, validate_route

# The report's Re-run Needed sheet: header on row 4, data from row 5,
# columns Airline (A), Route (B), Missing (D).
_RERUN_FIRST_ROW = 5


def configured_choices(
    commands: list[dict], domestic_airports=("DAC",)
) -> tuple[list[str], list[str]]:
    """Routes and airlines from configured commands, in report order:
    routes by destination with the domestic departure first (DAC-BKK, then
    BKK-DAC); airlines BS, BG, then alphabetical."""
    routes: set[str] = set()
    airlines: set[str] = set()
    for c in commands:
        origin = str(c.get("origin") or "").upper()
        dest = str(c.get("destination") or "").upper()
        if origin and dest:
            routes.add(f"{origin}-{dest}")
        if c.get("airline"):
            airlines.add(str(c["airline"]).upper())

    def route_key(route: str):
        origin, dest = route.split("-", 1)
        if origin in domestic_airports:
            return (dest, 0, origin)
        if dest in domestic_airports:
            return (origin, 1, dest)
        return (dest, 0, origin)

    return sorted(routes, key=route_key), sorted(airlines, key=_airline_sort_key)


def filter_problems(route_text: str, airline_text: str) -> list[str]:
    """Human-readable problems in the Route and Airline fields ([] = fine)."""
    problems = []
    for raw in (route_text or "").split(","):
        raw = raw.strip()
        if not raw:
            continue
        try:
            validate_route(raw)
        except ValidationError as exc:
            problems.append(f"Route '{raw}': {exc}")
    for raw in (airline_text or "").split(","):
        raw = raw.strip()
        if not raw:
            continue
        try:
            validate_airline_code(raw)
        except ValidationError as exc:
            problems.append(f"Airline '{raw}': {exc}")
    return problems


def rerun_pairs_from_report(path: str) -> list[dict]:
    """Rows of a fare report's Re-run Needed sheet ([] when it has none)."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True)
    except Exception:
        return []
    try:
        if RERUN_SHEET not in wb.sheetnames:
            return []
        pairs = []
        for row in wb[RERUN_SHEET].iter_rows(
            min_row=_RERUN_FIRST_ROW, max_col=4, values_only=True
        ):
            airline, route = (row[0] or ""), (row[1] or "")
            if airline and route:
                pairs.append(
                    {
                        "airline": str(airline).strip().upper(),
                        "route": str(route).strip().upper(),
                        "missing": str(row[3] or ""),
                    }
                )
        return pairs
    finally:
        wb.close()


def find_latest_rerun_report(reports_dir: str) -> str | None:
    """The newest full fare report, if it lists routes to re-run.

    Only the newest report counts: once a later run is clean, older lists
    are out of date.
    """
    reports = [
        p
        for p in glob.glob(os.path.join(reports_dir, "fare_report_*.xlsx"))
        if "_partial" not in os.path.basename(p)
        and not os.path.basename(p).startswith("~$")
    ]
    if not reports:
        return None
    newest = max(reports, key=os.path.getmtime)
    return newest if rerun_pairs_from_report(newest) else None


def pairs_to_text(pairs: list[dict]) -> str:
    """Pairs as the --pairs value: "BS:DAC-BKK,FZ:DXB-DAC"."""
    return ",".join(f"{p['airline']}:{p['route']}" for p in pairs)


def resumable_run(last_run: dict | None, checkpoint_dir: str) -> dict | None:
    """The last run's checkpoint if that run stopped before finishing.

    ``last_run`` is what the GUI saved when the run started: the checkpoint
    session name, the settings used and whether it finished. Returns
    {"path", "completed", "settings"} or None when there is nothing to resume.
    """
    import json

    if not last_run or last_run.get("finished") or not last_run.get("session"):
        return None
    path = os.path.join(checkpoint_dir, f"checkpoint_{last_run['session']}.json")
    try:
        with open(path, encoding="utf-8") as fh:
            completed = len(json.load(fh).get("completed_commands") or [])
    except (OSError, ValueError):
        return None
    if completed == 0:
        return None
    return {"path": path, "completed": completed, "settings": last_run.get("settings")}
