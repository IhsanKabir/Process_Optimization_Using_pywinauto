"""
baggage_report.py — Standalone Excel report for baggage allowance data.

Layout mirrors the YQ/YR charges report: grouped by route, one table per
destination, airlines as rows with Checked and Carry-on columns.
"""
from __future__ import annotations

import json
import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ── Styles ────────────────────────────────────────────────────────────────────
_TITLE_FONT = Font(name="Calibri", bold=True, size=16)
_HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_SECTION_FILL = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
_SECTION_FONT = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
_ALT_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
_BOLD = Font(name="Calibri", bold=True)
_NORMAL = Font(name="Calibri")
_NONE_FONT = Font(name="Calibri", italic=True, color="999999")

_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center")


def _c(ws, row: int, col: int, value, font=None, fill=None, align=None):
    cell = ws.cell(row=row, column=col, value=value)
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    return cell


def generate_baggage_report(
    baggage_data: dict,
    config: dict,
    output_path: str,
) -> str:
    """
    Write a standalone baggage allowance Excel report.

    Args:
        baggage_data: {file_key: {"checked": "30K", "carry_on": "07K", "route": "DACDXB"}}
        config: Loaded config dict (for city_names, airline_names).
        output_path: Destination .xlsx path.

    Returns:
        Absolute path to the written file.
    """
    airline_names: dict = config.get("airline_names", {})
    city_names: dict = config.get("city_names", {})
    domestic_airports: list = config.get("domestic_airports", [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Baggage Allowance"

    # ── Title ─────────────────────────────────────────────────────────────────
    row = 1
    _c(ws, row, 1, "Baggage Allowance Report", _TITLE_FONT)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 1
    _c(
        ws, row, 1,
        f"Generated: {datetime.now().strftime('%d-%b-%Y %H:%M')}",
        Font(name="Calibri", italic=True, size=10),
    )
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 2

    # ── Group by (destination, direction) ────────────────────────────────────
    # file_key format: "{AIRLINE}_{SRC}-{DST}"
    groups: dict[tuple, list] = {}
    for file_key, bag in baggage_data.items():
        parts = file_key.split("_", 1)
        if len(parts) != 2 or "-" not in parts[1]:
            continue
        airline = parts[0]
        route_str = parts[1]           # e.g. "DAC-DXB"
        src, dst = route_str.split("-", 1)
        is_domestic_src = src in domestic_airports
        is_domestic_dst = dst in domestic_airports
        if is_domestic_src:
            key = (dst, "outbound")
        elif is_domestic_dst:
            key = (src, "inbound")
        else:
            key = (dst, "outbound")
        groups.setdefault(key, []).append((airline, src, dst, bag))

    col_widths = [22, 14, 14, 18]

    for section_key in sorted(groups.keys()):
        intl_code, direction = section_key
        intl_name = city_names.get(intl_code, intl_code)
        arrow = "→" if direction == "outbound" else "←"

        # Section header
        _c(
            ws, row, 1,
            f"{arrow} {intl_name} ({intl_code})",
            _SECTION_FONT, _SECTION_FILL, _LEFT,
        )
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1

        # Column headers
        for col, label in enumerate(["Airline", "Checked", "Carry-on", "Route"], start=1):
            _c(ws, row, col, label, _HEADER_FONT, _HEADER_FILL, _CENTER)
        row += 1

        entries = groups[section_key]
        entries.sort(key=lambda e: e[0])
        for i, (airline, src, dst, bag) in enumerate(entries):
            fill = _ALT_FILL if i % 2 == 0 else None
            al_name = f"{airline_names.get(airline, airline)} ({airline})"
            checked = bag.get("checked") or "—"
            carry_on = bag.get("carry_on") or "—"
            route_label = f"{src}-{dst}"

            _c(ws, row, 1, al_name, _BOLD if not fill else Font(name="Calibri", bold=True), fill, _LEFT)
            _c(ws, row, 2, checked, _NORMAL, fill, _CENTER)
            _c(ws, row, 3, carry_on, _NORMAL, fill, _CENTER)
            _c(ws, row, 4, route_label, _NORMAL, fill, _CENTER)
            row += 1

        row += 1  # blank row between sections

    # ── Column widths ─────────────────────────────────────────────────────────
    for col, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "A5"

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    return os.path.abspath(output_path)


def save_baggage_json(baggage_data: dict, json_path: str) -> str:
    """Persist baggage_data dict to JSON for later use with --baggage-file."""
    os.makedirs(os.path.dirname(os.path.abspath(json_path)), exist_ok=True)
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(baggage_data, f, indent=2, ensure_ascii=False)
    return os.path.abspath(json_path)
