"""
currency_archive.py - Persistence layer for the Currency Rate report.

Writes today's rates to a per-date JSON file under data/archive/currency/,
reads the most recent prior-date snapshot for the "Previous Date" column,
and tracks the USD last-changed date for the special USD bracket annotation.

File layout
-----------
    data/archive/currency/
        rates_YYYY-MM-DD.json      per-date snapshot (latest run wins)
        usd_tracker.json           { "rate": 122.94, "last_changed_date": "2026-04-04" }

Importers accept JSON, CSV, and XLSX shapes so a user can replace the
previous-day baseline from the GUI or CLI.
"""

from __future__ import annotations

import csv
import json
import logging
import os
import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import Optional

logger = logging.getLogger("travelport.currency_archive")

ARCHIVE_DIR = os.path.join("data", "archive", "currency")
USD_TRACKER_FILE = "usd_tracker.json"
RATE_FILE_PATTERN = re.compile(r"^rates_(\d{4}-\d{2}-\d{2})\.json$")
DATE_FMT = "%Y-%m-%d"

# Extracts the leading numeric portion from a rate cell such as
# "122.94", "122.94 (04-APR-26)", or "122,940.50".
_RE_RATE_NUM = re.compile(r"-?\d+(?:\.\d+)?")
# Extracts a date bracket like "(04-APR-26)" or "(04 APR 2026)".
_RE_BRACKET_DATE = re.compile(
    r"\(\s*(\d{1,2})[-\s]+([A-Za-z]{3})[-\s]+(\d{2,4})\s*\)"
)
_MONTH_MAP = {
    "JAN": 1, "FEB": 2, "MAR": 3, "APR": 4, "MAY": 5, "JUN": 6,
    "JUL": 7, "AUG": 8, "SEP": 9, "OCT": 10, "NOV": 11, "DEC": 12,
}


def _parse_rate_value(value) -> Optional[float]:
    """Extract the numeric rate from a cell value, ignoring trailing annotations."""
    if value is None:
        return None
    s = str(value).replace(",", "").strip()
    if not s:
        return None
    match = _RE_RATE_NUM.search(s)
    if not match:
        return None
    try:
        return float(match.group(0))
    except ValueError:
        return None


def _parse_bracket_date(value) -> Optional[date]:
    """Extract a bracket date like '(04-APR-26)' from a cell value."""
    if value is None:
        return None
    match = _RE_BRACKET_DATE.search(str(value))
    if not match:
        return None
    day_s, mon_s, year_s = match.group(1), match.group(2).upper(), match.group(3)
    month = _MONTH_MAP.get(mon_s)
    if month is None:
        return None
    try:
        day = int(day_s)
        year = int(year_s)
        if year < 100:
            year += 2000
        return date(year, month, day)
    except ValueError:
        return None


@dataclass(frozen=True)
class RateSnapshot:
    snapshot_date: date
    rates: dict  # currency code (str, UPPER) -> rate (float)
    source: str = "live"


@dataclass(frozen=True)
class UsdTracker:
    rate: float
    last_changed_date: date


# ── Path helpers ────────────────────────────────────────────────────────────


def _archive_dir() -> str:
    os.makedirs(ARCHIVE_DIR, exist_ok=True)
    return ARCHIVE_DIR


def _rate_file(snapshot_date: date) -> str:
    return os.path.join(_archive_dir(), f"rates_{snapshot_date.strftime(DATE_FMT)}.json")


def _tracker_file() -> str:
    return os.path.join(_archive_dir(), USD_TRACKER_FILE)


# ── Snapshot I/O ────────────────────────────────────────────────────────────


def load_snapshot(snapshot_date: date) -> Optional[RateSnapshot]:
    path = _rate_file(snapshot_date)
    if not os.path.exists(path):
        return None
    try:
        with open(path, "r", encoding="utf-8") as fh:
            data = json.load(fh)
    except (OSError, json.JSONDecodeError) as exc:
        logger.warning("  Could not read %s: %s", path, exc)
        return None
    rates = {str(k).upper(): float(v) for k, v in (data.get("rates") or {}).items()}
    raw_date = data.get("date") or snapshot_date.strftime(DATE_FMT)
    try:
        parsed = datetime.strptime(raw_date, DATE_FMT).date()
    except ValueError:
        parsed = snapshot_date
    return RateSnapshot(snapshot_date=parsed, rates=rates, source=data.get("source", "live"))


def save_snapshot(snapshot: RateSnapshot) -> str:
    path = _rate_file(snapshot.snapshot_date)
    payload = {
        "date": snapshot.snapshot_date.strftime(DATE_FMT),
        "source": snapshot.source,
        "rates": {k.upper(): float(v) for k, v in snapshot.rates.items()},
    }
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, indent=2, sort_keys=True)
    return path


def latest_prior_snapshot(before_date: date) -> Optional[RateSnapshot]:
    """Return the most recent snapshot whose date is strictly earlier than before_date."""
    candidates: list[date] = []
    try:
        for name in os.listdir(_archive_dir()):
            match = RATE_FILE_PATTERN.match(name)
            if not match:
                continue
            try:
                d = datetime.strptime(match.group(1), DATE_FMT).date()
            except ValueError:
                continue
            if d < before_date:
                candidates.append(d)
    except FileNotFoundError:
        return None
    if not candidates:
        return None
    return load_snapshot(max(candidates))


# ── USD tracker ─────────────────────────────────────────────────────────────


def load_usd_tracker() -> Optional[UsdTracker]:
    path = _tracker_file()
    if not os.path.exists(path):
        return None
    try:
        with open(path, "r", encoding="utf-8") as fh:
            data = json.load(fh)
        return UsdTracker(
            rate=float(data["rate"]),
            last_changed_date=datetime.strptime(
                data["last_changed_date"], DATE_FMT
            ).date(),
        )
    except (OSError, KeyError, ValueError, json.JSONDecodeError) as exc:
        logger.warning("  Could not read USD tracker: %s", exc)
        return None


def save_usd_tracker(tracker: UsdTracker) -> None:
    path = _tracker_file()
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(
            {
                "rate": tracker.rate,
                "last_changed_date": tracker.last_changed_date.strftime(DATE_FMT),
            },
            fh,
            indent=2,
        )


def update_usd_tracker(current_rate: float, run_date: date) -> UsdTracker:
    """Compare current USD rate against tracker; update last_changed_date if changed."""
    existing = load_usd_tracker()
    if existing is None:
        tracker = UsdTracker(rate=current_rate, last_changed_date=run_date)
        save_usd_tracker(tracker)
        return tracker
    if not _rates_equal(existing.rate, current_rate):
        tracker = UsdTracker(rate=current_rate, last_changed_date=run_date)
        save_usd_tracker(tracker)
        return tracker
    return existing


def _rates_equal(a: float, b: float, tol: float = 1e-6) -> bool:
    return abs(a - b) <= tol


# ── Importers (user-supplied previous data) ─────────────────────────────────


def import_previous_rates(path: str, snapshot_date: date) -> RateSnapshot:
    """Load a user-supplied file and persist it as the snapshot for snapshot_date.

    Accepts .json, .csv, .xlsx. Unknown extensions are rejected.

    When the USD row carries a bracket date like '(04-APR-26)', that date is
    also written to the USD tracker so the current-day USD cell can render
    the correct "last changed" annotation.
    """
    ext = os.path.splitext(path)[1].lower()
    usd_last_changed: Optional[date] = None
    if ext == ".json":
        rates = _read_json_rates(path)
    elif ext == ".csv":
        rates, usd_last_changed = _read_csv_rates(path)
    elif ext in (".xlsx", ".xlsm"):
        rates, usd_last_changed = _read_xlsx_rates(path)
    else:
        raise ValueError(f"Unsupported file type: {ext} (expected .json, .csv, .xlsx)")

    if not rates:
        raise ValueError(f"No valid rates found in {path}")

    snapshot = RateSnapshot(snapshot_date=snapshot_date, rates=rates, source="imported")
    save_snapshot(snapshot)

    usd_rate = rates.get("USD")
    if usd_rate and usd_last_changed:
        save_usd_tracker(UsdTracker(rate=usd_rate, last_changed_date=usd_last_changed))
        logger.info(
            "  [IMPORT] Seeded USD tracker: rate=%s, last_changed=%s",
            usd_rate, usd_last_changed.strftime(DATE_FMT),
        )
    return snapshot


def _read_json_rates(path: str) -> dict:
    with open(path, "r", encoding="utf-8") as fh:
        data = json.load(fh)
    if isinstance(data, dict) and "rates" in data and isinstance(data["rates"], dict):
        return _coerce_rate_map(data["rates"])
    if isinstance(data, dict):
        return _coerce_rate_map(data)
    if isinstance(data, list):
        flat = {}
        for row in data:
            if isinstance(row, dict) and "currency" in row and "rate" in row:
                flat[str(row["currency"]).upper()] = float(row["rate"])
        return flat
    return {}


def _read_csv_rates(path: str) -> tuple[dict, Optional[date]]:
    rates: dict = {}
    usd_last_changed: Optional[date] = None
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        rows = list(reader)
    if not rows:
        return rates, None
    header = [c.strip().lower() for c in rows[0]]
    if "currency" in header and "rate" in header:
        cur_i = header.index("currency")
        rate_i = header.index("rate")
        data_rows = rows[1:]
    else:
        cur_i, rate_i = 0, 1
        data_rows = rows
    for row in data_rows:
        if len(row) <= max(cur_i, rate_i):
            continue
        cur_raw = row[cur_i].strip().upper()
        # Accept "1 USD" or "USD"
        cur = cur_raw.split()[-1] if cur_raw else ""
        if len(cur) != 3 or not cur.isalpha():
            continue
        rate_val = _parse_rate_value(row[rate_i])
        if rate_val is None:
            continue
        rates[cur] = rate_val
        if cur == "USD" and usd_last_changed is None:
            usd_last_changed = _parse_bracket_date(row[rate_i])
    return rates, usd_last_changed


def _read_xlsx_rates(path: str) -> tuple[dict, Optional[date]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl required to import XLSX rates") from exc

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rates: dict = {}
    usd_last_changed: Optional[date] = None
    rows_iter = ws.iter_rows(values_only=True)
    try:
        first = next(rows_iter)
    except StopIteration:
        return rates, None

    def _is_header(cells) -> bool:
        return any(
            isinstance(v, str) and v.strip().lower() in ("currency", "rate", "exchange rate to bdt")
            for v in cells
            if v is not None
        )

    data_rows = list(rows_iter) if _is_header(first) else [first, *rows_iter]
    for row in data_rows:
        if not row:
            continue
        cur_cell = row[0]
        rate_cell = row[1] if len(row) > 1 else None
        if cur_cell is None or rate_cell is None:
            continue
        cur_raw = str(cur_cell).strip().upper()
        cur = cur_raw.split()[-1] if cur_raw else ""
        if len(cur) != 3 or not cur.isalpha():
            continue
        rate_val = _parse_rate_value(rate_cell)
        if rate_val is None:
            continue
        rates[cur] = rate_val
        if cur == "USD" and usd_last_changed is None:
            usd_last_changed = _parse_bracket_date(rate_cell)
    return rates, usd_last_changed


def _coerce_rate_map(raw: dict) -> dict:
    out: dict = {}
    for k, v in raw.items():
        key = str(k).strip().upper()
        key = key.split()[-1] if key else ""
        if len(key) != 3 or not key.isalpha():
            continue
        try:
            out[key] = float(v)
        except (TypeError, ValueError):
            continue
    return out
