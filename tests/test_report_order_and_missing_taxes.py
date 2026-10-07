"""Fare report: route/airline ordering, missing-tax flags and original-currency
tax amounts."""

from openpyxl import Workbook, load_workbook

from excel_report import (
    _airline_sort_key,
    _collect_missing_tax_routes,
    _section_sort_key,
    _write_individual_tables_sheet,
    _write_tax_breakdown_sheet,
    generate_report,
)

TAXES = {
    "base_currency": "USD",
    "equ_currency": "BDT",
    "exchange_rate": 120.0,
    "roe": 1.0,
    "yq_charge": 1200.0,
    "yr_charge": 0.0,
    "q_charge": 10.0,
    "total_taxes": 6000.0,
    "total_amount": 42000.0,
    "tax_breakdown": {"BD": 3000.0, "E5": 600.0},
}
FARES = {"Y": {"ow_fare": 300, "rt_fare": 550}}


def _route(fs_taxes=None, fares=FARES):
    return {"currency": "USD", "fs_taxes": dict(fs_taxes or {}), "rbd_data": fares}


def _titles(ws, row):
    return [
        ws.cell(row=row, column=c).value
        for c in range(1, ws.max_column + 1)
        if ws.cell(row=row, column=c).value
    ]


def _header_row(ws, col=1):
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, col).value == "RBD")


# ── ordering ────────────────────────────────────────────────────────────────


def test_outbound_section_comes_before_inbound():
    keys = [("BKK", "inbound"), ("BKK", "outbound"), ("CCU", "inbound")]

    assert sorted(keys, key=_section_sort_key) == [
        ("BKK", "outbound"),
        ("BKK", "inbound"),
        ("CCU", "inbound"),
    ]


def test_bs_then_bg_then_alphabetical():
    airlines = ["UL", "BG", "8D", "BS", "EK"]

    assert sorted(airlines, key=_airline_sort_key) == ["BS", "BG", "8D", "EK", "UL"]


def test_main_report_lists_dac_bkk_before_bkk_dac_and_bs_first(tmp_path):
    data = {
        "BG_BKK-DAC": _route(TAXES),
        "BS_BKK-DAC": _route(TAXES),
        "BG_DAC-BKK": _route(TAXES),
        "BS_DAC-BKK": _route(TAXES),
    }
    out = str(tmp_path / "fares.xlsx")

    generate_report(data, out, config={"domestic_airports": ["DAC"]})

    ws = load_workbook(out)["Individual Tables"]
    titles = [
        ws.cell(r, c).value
        for r in range(1, ws.max_row + 1)
        for c in range(1, ws.max_column + 1)
        if isinstance(ws.cell(r, c).value, str) and " / " in ws.cell(r, c).value
    ]
    assert [t.split()[0:3] for t in titles] == [
        ["BS", "/", "DAC-BKK"],
        ["BG", "/", "DAC-BKK"],
        ["BS", "/", "BKK-DAC"],
        ["BG", "/", "BKK-DAC"],
    ]


# ── missing taxes ───────────────────────────────────────────────────────────


def _individual(all_route_data, route_key):
    wb = Workbook()
    ws = wb.active
    airline, route = route_key.split("_")
    origin, dest = route.split("-")
    intl, direction = (dest, "outbound") if origin == "DAC" else (origin, "inbound")
    sections = {
        (intl, direction): [(airline, "DAC", route_key, all_route_data[route_key])]
    }
    _write_individual_tables_sheet(
        ws, all_route_data, sections, {}, {}, [], ["DAC"], None
    )
    return ws


def test_missing_taxes_flag_title_and_keep_only_fare_columns():
    data = {"BS_DAC-BKK": _route({}), "BS_BKK-DAC": _route(TAXES)}

    ws = _individual(data, "BS_DAC-BKK")

    hr = _header_row(ws)
    assert _titles(ws, hr) == ["RBD", "OW/USD", "RT/USD"]
    title = ws.cell(hr - 2, 1)
    assert "TAXES MISSING" in title.value
    assert title.fill.start_color.rgb.endswith("FFC7CE")


def test_missing_return_leg_taxes_drops_only_rt_gross():
    data = {"BS_DAC-BKK": _route(TAXES), "BS_BKK-DAC": _route({})}

    ws = _individual(data, "BS_DAC-BKK")

    hr = _header_row(ws)
    assert _titles(ws, hr) == [
        "RBD",
        "OW/USD",
        "With YQ/OW(USD)",
        "OW/Gross(BDT)",
        "RT/USD",
        "With YQ/RT(USD)",
    ]
    assert "return" in ws.cell(hr - 2, 1).value.lower()


def test_complete_taxes_keep_all_columns_without_flag():
    data = {"BS_DAC-BKK": _route(TAXES), "BS_BKK-DAC": _route(TAXES)}

    ws = _individual(data, "BS_DAC-BKK")

    hr = _header_row(ws)
    assert _titles(ws, hr)[-1] == "RT/Gross(BDT)"
    assert "MISSING" not in ws.cell(hr - 2, 1).value


def test_collect_missing_tax_routes_skips_routes_without_fares():
    data = {
        "BG_BKK-DAC": _route({}),
        "BS_DAC-BKK": _route({}),
        "UL_DAC-CMB": _route({}, fares={}),
        "BS_BKK-DAC": _route(TAXES),
    }

    missing = _collect_missing_tax_routes(data, ["DAC"])

    assert [(m["airline"], m["route"]) for m in missing] == [
        ("BS", "DAC-BKK"),
        ("BG", "BKK-DAC"),
    ]


def test_report_has_rerun_sheet_only_when_taxes_missing(tmp_path):
    out = str(tmp_path / "fares.xlsx")
    generate_report(
        {"BS_DAC-BKK": _route({}), "BS_BKK-DAC": _route(TAXES)},
        out,
        config={"domestic_airports": ["DAC"]},
    )

    wb = load_workbook(out)
    assert wb.sheetnames[1] == "Re-run Needed"
    rows = [r for r in wb["Re-run Needed"].iter_rows(values_only=True) if any(r)]
    assert ("BS", "DAC-BKK") in [(r[0], r[1]) for r in rows]

    clean = str(tmp_path / "clean.xlsx")
    generate_report(
        {"BS_DAC-BKK": _route(TAXES), "BS_BKK-DAC": _route(TAXES)},
        clean,
        config={"domestic_airports": ["DAC"]},
    )
    assert "Re-run Needed" not in load_workbook(clean).sheetnames


def test_fd_only_runs_do_not_flag_missing_taxes(tmp_path):
    out = str(tmp_path / "fd_only.xlsx")

    generate_report(
        {"BS_DAC-BKK": _route({}), "BS_BKK-DAC": _route({})},
        out,
        config={"domestic_airports": ["DAC"]},
        taxes_expected=False,
    )

    wb = load_workbook(out)
    assert "Re-run Needed" not in wb.sheetnames
    texts = [
        v
        for row in wb["Individual Tables"].iter_rows(values_only=True)
        for v in row
        if isinstance(v, str)
    ]
    assert not any("MISSING" in t for t in texts)


# ── original currency on Tax Breakdowns ─────────────────────────────────────


def test_tax_breakdown_shows_original_currency_beside_bdt():
    wb = Workbook()
    ws = wb.active
    data = {"BS_DAC-BKK": _route(TAXES)}
    sections = {("BKK", "outbound"): [("BS", "DAC", "BS_DAC-BKK", data["BS_DAC-BKK"])]}

    _write_tax_breakdown_sheet(ws, data, sections, {}, {}, ["DAC"])

    rows = {
        ws.cell(r, 1).value: (ws.cell(r, 2).value, ws.cell(r, 3).value)
        for r in range(1, ws.max_row + 1)
    }
    header_row = next(
        r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Tax/Charge"
    )
    assert ws.cell(header_row, 2).value == "Amount (BDT)"
    assert ws.cell(header_row, 3).value == "Original (USD)"
    assert rows["BD"] == (3000.0, 25.0)
    assert rows["YQ"] == (1200.0, 10.0)
    assert rows["Q"][1] == 10.0  # NUC 10 x ROE 1
    assert rows["Total Taxes"] == (6000.0, 50.0)
