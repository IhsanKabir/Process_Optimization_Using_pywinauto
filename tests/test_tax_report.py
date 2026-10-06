"""Tests for tax_report.generate_tax_report workbook layout."""

import openpyxl
import pytest

from tax_report import generate_tax_report

L7_DATA = {
    "SIN": {
        "_country": "SINGAPORE",
        "taxes": [
            {
                "code": "L7",
                "name": "AIRPORT DEVELOPMENT LEVY",
                "sections": [
                    {
                        "category": "DEPARTURES FROM CHANGI SIN",
                        "subcategory": "TERMINAL 1 2 3",
                        "rates": [
                            {
                                "condition": "TKT ON/BEFORE 31MAR25",
                                "currency": "SGD",
                                "amount": 46.40,
                                "status": "expired",
                            },
                            {
                                "condition": "TVL ON/AFTER 01APR25",
                                "currency": "SGD",
                                "amount": 49.40,
                                "status": "current",
                            },
                        ],
                    }
                ],
                "exemptions": [],
            }
        ],
    }
}


def _rows(path: str, sheet: str) -> list[tuple]:
    ws = openpyxl.load_workbook(path)[sheet]
    return [row for row in ws.iter_rows(values_only=True)]


@pytest.fixture
def out(tmp_path):
    return str(tmp_path / "reports" / "tax.xlsx")


def test_creates_summary_and_details_sheets_without_changes(out):
    generate_tax_report(L7_DATA, out)

    assert openpyxl.load_workbook(out).sheetnames == ["Tax Summary", "Detailed Rates"]


def test_summary_has_one_row_per_rate(out):
    generate_tax_report(L7_DATA, out)

    rows = _rows(out, "Tax Summary")
    assert rows[0][0] == "Airport"
    assert rows[1] == (
        "SIN",
        "SINGAPORE",
        "L7",
        "AIRPORT DEVELOPMENT LEVY",
        "DEPARTURES FROM CHANGI SIN TERMINAL 1 2 3",
        "TKT ON/BEFORE 31MAR25",
        "SGD",
        46.40,
        "Expired",
    )
    assert rows[2][-1] == "Current ●"
    assert len(rows) == 3


def test_country_name_from_config_overrides_data(out):
    config = {"tax_airports": {"SIN": {"name": "Singapore Changi"}}}

    generate_tax_report(L7_DATA, out, config=config)

    assert _rows(out, "Tax Summary")[1][1] == "Singapore Changi"


def test_percent_rate_rendered_with_basis_codes(out):
    data = {
        "DAC": {
            "taxes": [
                {
                    "code": "E5",
                    "name": "VAT",
                    "sections": [
                        {
                            "category": "",
                            "subcategory": "",
                            "rates": [
                                {
                                    "condition": "",
                                    "currency": None,
                                    "amount": None,
                                    "percent": 15.0,
                                    "basis_codes": ["BD", "P7"],
                                    "status": "current",
                                }
                            ],
                        }
                    ],
                    "exemptions": [
                        {
                            "pax_type": "INFANTS",
                            "rates": [
                                {
                                    "condition": "",
                                    "percent": 15.0,
                                    "basis_codes": ["P7"],
                                    "status": "current",
                                }
                            ],
                        }
                    ],
                }
            ]
        }
    }

    generate_tax_report(data, out)

    rows = _rows(out, "Tax Summary")
    assert rows[1][7] == "15% of BD+P7"
    assert rows[2][4] == "[INFANTS]"
    assert rows[2][7] == "15% of P7"


def test_tax_without_rates_still_gets_a_row(out):
    data = {
        "MLE": {
            "taxes": [
                {"code": "YY", "name": "Q2 TAX", "sections": [], "exemptions": []},
                {
                    "code": "BQ",
                    "name": "BQ TAX",
                    "sections": [
                        {
                            "category": "",
                            "subcategory": "",
                            "rates": [
                                {
                                    "condition": "",
                                    "currency": "USD",
                                    "amount": 30.0,
                                    "status": "current",
                                }
                            ],
                        }
                    ],
                    "exemptions": [],
                },
            ]
        }
    }

    generate_tax_report(data, out)

    rows = _rows(out, "Tax Summary")
    assert [r[2] for r in rows[1:]] == ["YY", "BQ"]
    assert rows[1][5:] == ("—", "—", "—", "—")


def test_details_sheet_has_title_and_rates(out):
    generate_tax_report(L7_DATA, out)

    rows = _rows(out, "Detailed Rates")
    assert rows[0][0] == "SINGAPORE (SIN) — AIRPORT DEVELOPMENT LEVY (L7)"
    assert rows[1][:4] == ("Condition", "Currency", "Amount", "Status")
    assert rows[3][:4] == ("TKT ON/BEFORE 31MAR25", "SGD", 46.40, "Expired")


def test_changes_sheet_lists_each_change(out):
    changes = {
        "SIN": {
            "L7": [
                {
                    "type": "amount_changed",
                    "section": "TERMINAL 1 2 3",
                    "condition": "TVL ON/AFTER 01APR25",
                    "old_amount": 46.40,
                    "new_amount": 49.40,
                    "currency": "SGD",
                    "status": "current",
                }
            ]
        }
    }

    generate_tax_report(L7_DATA, out, changes=changes)

    rows = _rows(out, "Changes")
    assert rows[0][0] == "Tax Changes Summary"
    assert rows[4] == (
        "SIN",
        "L7",
        "AIRPORT DEVELOPMENT LEVY",
        "TERMINAL 1 2 3",
        "TVL ON/AFTER 01APR25",
        "AMOUNT_CHANGED",
        46.40,
        49.40,
        "SGD",
        "Current",
    )


def test_empty_changes_dict_skips_changes_sheet(out):
    generate_tax_report(L7_DATA, out, changes={"SIN": {}})

    assert "Changes" not in openpyxl.load_workbook(out).sheetnames
