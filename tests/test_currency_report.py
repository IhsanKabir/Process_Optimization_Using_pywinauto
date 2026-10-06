"""Tests for currency_report and its currency_archive persistence."""

from datetime import date

import openpyxl
import pytest

import currency_archive
from currency_archive import RateSnapshot, UsdTracker, load_usd_tracker, save_snapshot
from currency_report import _format_rate, generate_currency_report

DAY1 = date(2026, 3, 1)
DAY2 = date(2026, 3, 2)
CHANGED_RGB = "00C6E7F5"


@pytest.fixture(autouse=True)
def archive_dir(tmp_path, monkeypatch):
    path = tmp_path / "archive"
    monkeypatch.setattr(currency_archive, "ARCHIVE_DIR", str(path))
    return path


@pytest.fixture
def out(tmp_path):
    return str(tmp_path / "out" / "currency.xlsx")


def _sheet(path: str):
    return openpyxl.load_workbook(path)["Currency Rate"]


def test_rows_sorted_by_rate_descending(out):
    generate_currency_report({"usd": 122.0, "KWD": 398.4, "AED": 33.2}, out, DAY1)

    ws = _sheet(out)
    assert [ws.cell(row=r, column=1).value for r in (4, 5, 6)] == [
        "1 KWD",
        "1 USD",
        "1 AED",
    ]


def test_zero_or_missing_rates_are_dropped(out):
    generate_currency_report({"KWD": 398.4, "XXX": 0, "YYY": None}, out, DAY1)

    ws = _sheet(out)
    assert ws.cell(row=4, column=1).value == "1 KWD"
    assert ws.cell(row=5, column=1).value is None


def test_first_run_has_no_previous_and_highlights_everything(out):
    result = generate_currency_report({"KWD": 398.4}, out, DAY1)

    ws = _sheet(out)
    assert result.previous_date is None
    assert ws.cell(row=2, column=4).value == "—"
    assert ws.cell(row=4, column=5).value == "—"
    assert ws.cell(row=4, column=2).fill.start_color.rgb == CHANGED_RGB


def test_previous_snapshot_used_and_unchanged_rate_not_highlighted(out):
    generate_currency_report({"KWD": 398.4, "AED": 33.2}, out, DAY1)

    result = generate_currency_report({"KWD": 398.4, "AED": 33.5}, out, DAY2)

    ws = _sheet(out)
    assert result.previous_date == DAY1
    assert ws.cell(row=2, column=4).value == "01-Mar-26"
    kwd_row, aed_row = 4, 5
    assert ws.cell(row=kwd_row, column=5).value == 398.4
    assert ws.cell(row=kwd_row, column=2).fill.fill_type is None
    assert ws.cell(row=aed_row, column=2).fill.start_color.rgb == CHANGED_RGB


def test_zenith_is_inverse_rate(out):
    generate_currency_report({"KWD": 400.0}, out, DAY1)

    assert _sheet(out).cell(row=4, column=8).value == pytest.approx(0.0025)


def test_usd_cell_shows_last_changed_date(out):
    generate_currency_report({"USD": 122.0}, out, DAY1)
    generate_currency_report({"USD": 122.0}, out, DAY2)

    value = _sheet(out).cell(row=4, column=2).value
    assert "122.00" in str(value)
    assert "(01-MAR-26)" in str(value)


def test_usd_tracker_moves_when_rate_changes(out):
    generate_currency_report({"USD": 122.0}, out, DAY1)

    result = generate_currency_report({"USD": 123.0}, out, DAY2)

    assert result.usd_tracker == UsdTracker(rate=123.0, last_changed_date=DAY2)
    assert load_usd_tracker() == result.usd_tracker


def test_persist_false_writes_nothing_to_archive(out, archive_dir):
    generate_currency_report({"USD": 122.0}, out, DAY1, persist=False)

    assert not any(archive_dir.glob("rates_*.json"))
    assert load_usd_tracker() is None


def test_snapshot_round_trip_and_prior_lookup():
    save_snapshot(RateSnapshot(snapshot_date=DAY1, rates={"usd": 1.5}))

    assert currency_archive.load_snapshot(DAY1).rates == {"USD": 1.5}
    assert currency_archive.latest_prior_snapshot(DAY2).snapshot_date == DAY1
    assert currency_archive.latest_prior_snapshot(DAY1) is None


def test_corrupt_snapshot_returns_none(archive_dir):
    archive_dir.mkdir(parents=True)
    (archive_dir / f"rates_{DAY1.isoformat()}.json").write_text("{bad")

    assert currency_archive.load_snapshot(DAY1) is None


@pytest.mark.parametrize(
    "rate, expected",
    [(122.0, "122.00"), (122.5, "122.50"), (398.427041, "398.427041"), (1.25, "1.25")],
)
def test_format_rate(rate, expected):
    assert _format_rate(rate) == expected
