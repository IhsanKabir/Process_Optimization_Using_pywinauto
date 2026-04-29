import gc
from pathlib import Path

import openpyxl

from penalty_parser import extract_penalty_section, parse_penalty_text
from penalty_report import (
    _bucket_records_into_comparison_rows,
    _classify_rule_bucket,
    _format_rule_cell,
    generate_penalty_report,
)

RAW_SAMPLE = """
FARES LAST UPDATED 09APR  1:29 A
GF       DAC-MED DEPART 09APR
MPM 3812 EH
**ADDITIONAL TAXES/FEES MAY APPLY**
PUBLIC FARES
     CX    FARE   FARE     C  AP  MIN/    SEASONS...... MR GI DT
           USD     BASIS             MAX
DACMED
  1  GF  460.00R  WCLIT1BD W       3/3M                 R  EH
  2  GF  480.00R  SCLIT1BD S       3/3M                 R  EH

                                                    Close
 16. PENALTIES
UNLESS OTHERWISE SPECIFIED   NOTE - RULE 0100 IN IPRG100
ALSO APPLIES
  CHANGES
    ANY TIME
      CHARGE USD 95.00 FOR REISSUE.
         NOTE -
          CHANGE FEES MUST BE CODED AS OA IN THE TAX FIELD.
    CHARGE USD 190.00 FOR NO-SHOW.
         NOTE -
          PASSENGER WILL BE TREATED AS NO SHOW IF THE
          BOOKING IS NOT CANCELLED AT LEAST 4 HOURS BEFORE
          THE DEPARTURE OF THE FLIGHT.
  CANCELLATIONS
    ANY TIME
      CHARGE USD 150.00 FOR CANCEL/REFUND.
    CHARGE USD 300.00 FOR NO-SHOW.
         NOTE -
          SERVICE CHARGES COLLECTED UNDER CODES
          DU/DV/OC/OA/OB/OD/OE ARE NON-REFUNDABLE.
UNLESS OTHERWISE SPECIFIED
  FOR TICKETING ON/AFTER 26MAR 26
    IF TRAVEL COMMENCES ON/BEFORE 31MAR 27
      FOR TRAVEL ON/BEFORE 31MAR 27
        CHANGES
          CHANGES PERMITTED.
         NOTE -
          FIRST CHANGE PERMITTED FOC UPTO FOUR HOURS BEFORE
          FLIGHT DEPARTURE AND SUBSEQUENT CHANGE AT CHARGE.

 14  GF  324.00   WDSMR3BD W                            R  EH
      END
"""


def test_extract_penalty_section_stops_before_fare_grid_resumes():
    section = extract_penalty_section(RAW_SAMPLE)

    assert section.startswith("16. PENALTIES")
    assert "CHARGE USD 95.00 FOR REISSUE." in section
    assert "14  GF  324.00" not in section


def test_parse_penalty_text_extracts_structured_rules():
    record = parse_penalty_text(
        RAW_SAMPLE,
        airline="GF",
        route="DAC-MED",
        fare_basis="WCLIT1BD",
        rbd="W",
        fare_amount=460.0,
        journey_type="RT",
    )

    assert record["fare_basis"] == "WCLIT1BD"
    assert len(record["rules"]) == 5

    change_rule = next(
        rule
        for rule in record["rules"]
        if rule["category"] == "CHANGES" and rule["amount"] == 95.0
    )
    assert change_rule["currency"] == "USD"
    assert change_rule["criteria_text"] == "ANY TIME"
    assert "CHANGE FEES MUST BE CODED AS OA" in change_rule["note_text"]

    no_show_refund_rule = next(
        rule
        for rule in record["rules"]
        if rule["category"] == "CANCELLATIONS" and rule["amount"] == 300.0
    )
    assert no_show_refund_rule["subtype"] == "no_show"
    assert "NON-REFUNDABLE" in no_show_refund_rule["note_text"]

    no_show_change_rule = next(
        rule
        for rule in record["rules"]
        if rule["category"] == "CHANGES" and rule["amount"] == 190.0
    )
    assert no_show_change_rule["timing_value"] == 4
    assert no_show_change_rule["timing_unit"] == "hour"
    assert no_show_change_rule["timing_direction"] == "before"
    assert no_show_change_rule["timing_reference"] == "departure"
    assert no_show_change_rule["timing_qualifier"] == "at_least"
    assert (
        "4 HOURS BEFORE THE DEPARTURE OF THE FLIGHT"
        in no_show_change_rule["timing_text"].upper()
    )

    permitted_rule = next(
        rule for rule in record["rules"] if rule["status"] == "permitted"
    )
    assert "FOR TICKETING ON/AFTER 26MAR 26" in permitted_rule["criteria_text"]
    assert "FOR TRAVEL ON/BEFORE 31MAR 27" in permitted_rule["criteria_text"]
    assert permitted_rule["timing_value"] == 4
    assert permitted_rule["timing_unit"] == "hour"
    assert permitted_rule["timing_direction"] == "before"
    assert permitted_rule["timing_reference"] == "departure"
    assert permitted_rule["timing_qualifier"] == "upto"


def test_generate_penalty_report_creates_expected_sheets():
    record = parse_penalty_text(
        RAW_SAMPLE,
        airline="GF",
        route="DAC-MED",
        fare_basis="WCLIT1BD",
        rbd="W",
        fare_amount=460.0,
        journey_type="RT",
    )

    output_path = Path.cwd() / "penalty_report_test.xlsx"
    workbook = None
    try:
        generate_penalty_report([record], str(output_path))

        workbook = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
        assert workbook.sheetnames == [
            "Comparison",
            "Penalty Summary",
            "Penalty Details",
            "Raw Captures",
        ]
        assert workbook["Penalty Summary"]["A2"].value == "GF"
        assert workbook["Penalty Details"]["D2"].value == "WCLIT1BD"
        assert workbook["Penalty Details"]["J1"].value == "Timing Text"
        detail_rows = list(
            workbook["Penalty Details"].iter_rows(min_row=2, values_only=True)
        )
        no_show_change_row = next(row for row in detail_rows if row[15] == 190)
        assert no_show_change_row[9]
        assert no_show_change_row[10] == "at_least"
        assert no_show_change_row[11] == 4
        assert no_show_change_row[12] == "hour"
        assert no_show_change_row[13] == "before"
    finally:
        if workbook is not None:
            workbook.close()
            workbook = None
            gc.collect()
        if output_path.exists():
            output_path.unlink()


# ── v1.5.16: side-by-side comparison sheet ─────────────────────────────────────


def _make_rule(
    category: str,
    *,
    subtype: str = "change",
    amount: float | None = None,
    currency: str | None = None,
    description: str = "",
    timing_qualifier: str | None = None,
    timing_unit: str | None = None,
    timing_direction: str | None = None,
    note_text: str = "",
) -> dict:
    return {
        "category": category,
        "subtype": subtype,
        "amount": amount,
        "currency": currency,
        "description": description,
        "timing_qualifier": timing_qualifier,
        "timing_unit": timing_unit,
        "timing_direction": timing_direction,
        "timing_reference": "departure",
        "timing_text": "",
        "criteria_text": "",
        "note_text": note_text,
        "status": "charged",
    }


def test_classify_rule_bucket_no_show_changes_to_reissue_noshow():
    rule = _make_rule("CHANGES", subtype="no_show", amount=190, currency="USD")
    assert _classify_rule_bucket(rule) == "reissue_noshow"


def test_classify_rule_bucket_no_show_cancellations_to_refund_noshow():
    rule = _make_rule("CANCELLATIONS", subtype="no_show", amount=300, currency="USD")
    assert _classify_rule_bucket(rule) == "refund_noshow"


def test_classify_rule_bucket_within_24h_change_to_reissue_within():
    rule = _make_rule(
        "CHANGES",
        subtype="change",
        amount=20,
        currency="USD",
        timing_qualifier="within",
        timing_unit="hour",
        timing_direction="before",
    )
    assert _classify_rule_bucket(rule) == "reissue_within_24h"


def test_classify_rule_bucket_general_change_defaults_to_before_24h():
    """A 'CHARGE USD 95.00 FOR REISSUE.' under ANY TIME has no timing
    details and must default to the before-24-hours bucket."""
    rule = _make_rule("CHANGES", subtype="change", amount=95, currency="USD")
    assert _classify_rule_bucket(rule) == "reissue_before_24h"


def test_classify_rule_bucket_unknown_category_returns_none():
    rule = _make_rule("MISC", subtype="other")
    assert _classify_rule_bucket(rule) is None


def test_format_rule_cell_renders_amount_and_currency():
    rule = _make_rule("CHANGES", amount=10, currency="USD")
    assert _format_rule_cell(rule) == "USD 10"


def test_format_rule_cell_appends_first_change_free_marker():
    rule = _make_rule(
        "CHANGES",
        amount=10,
        currency="USD",
        note_text="FIRST CHANGE PERMITTED FOC UPTO FOUR HOURS BEFORE FLIGHT DEPARTURE",
    )
    assert _format_rule_cell(rule).endswith("*")


def test_format_rule_cell_falls_back_to_description_when_no_amount():
    rule = _make_rule(
        "CHANGES",
        amount=None,
        currency=None,
        description="100% of Base Fare.",
    )
    assert _format_rule_cell(rule) == "100% of Base Fare"


def test_bucket_records_groups_same_buckets_under_one_rbds_row():
    """Two records on the same route/airline that produce identical bucket
    values across all six categories must collapse into a single row whose
    RBDs label combines both fare bases (e.g. 'K/B/H')."""
    rec_k = {
        "airline": "BG",
        "route": "DAC-CCU",
        "rbd": "K",
        "currency": "USD",
        "rules": [
            _make_rule("CHANGES", amount=10, currency="USD"),
            _make_rule(
                "CANCELLATIONS", subtype="refund", amount=15, currency="USD"
            ),
        ],
    }
    rec_b = {**rec_k, "rbd": "B"}
    rec_h = {**rec_k, "rbd": "H"}

    rows = _bucket_records_into_comparison_rows([rec_k, rec_b, rec_h])

    assert len(rows) == 1
    assert rows[0]["airline"] == "BG"
    assert set(rows[0]["rbds"].split("/")) == {"K", "B", "H"}
    assert rows[0]["buckets"]["reissue_before_24h"] == "USD 10"
    assert rows[0]["buckets"]["refund_before_24h"] == "USD 15"


def test_bucket_records_derives_currency_from_rules_when_record_has_none():
    """parse_penalty_text doesn't set a record-level currency field; the
    comparison sheet must fall back to the first rule's currency so the
    route header still displays 'DAC-CCU (USD)' style."""
    rec = {
        "airline": "BS",
        "route": "DAC-CCU",
        "rbd": "ALL",
        "rules": [
            _make_rule("CHANGES", amount=10, currency="USD"),
        ],
    }

    rows = _bucket_records_into_comparison_rows([rec])

    assert len(rows) == 1
    assert rows[0]["currency"] == "USD"


def test_bucket_records_keeps_rbd_groups_separate_when_buckets_differ():
    """Two RBDs with different penalty values must NOT collapse into one row."""
    rec_k = {
        "airline": "BG",
        "route": "DAC-CCU",
        "rbd": "K",
        "currency": "USD",
        "rules": [_make_rule("CHANGES", amount=10, currency="USD")],
    }
    rec_l = {
        "airline": "BG",
        "route": "DAC-CCU",
        "rbd": "L",
        "currency": "USD",
        "rules": [_make_rule("CHANGES", amount=20, currency="USD")],
    }

    rows = _bucket_records_into_comparison_rows([rec_k, rec_l])

    assert len(rows) == 2
    rbds = sorted(r["rbds"] for r in rows)
    assert rbds == ["K", "L"]


def test_comparison_sheet_renders_headers_and_one_row():
    rec = {
        "airline": "BS",
        "route": "DAC-CCU",
        "rbd": "ALL",
        "currency": "USD",
        "rules": [
            _make_rule("CHANGES", amount=10, currency="USD"),
            _make_rule(
                "CHANGES",
                subtype="no_show",
                amount=30,
                currency="USD",
                description="No show",
            ),
            _make_rule(
                "CANCELLATIONS", subtype="refund", amount=20, currency="USD"
            ),
            _make_rule(
                "CANCELLATIONS",
                subtype="no_show",
                amount=50,
                currency="USD",
                description="No show",
            ),
        ],
    }

    output_path = Path.cwd() / "penalty_comparison_test.xlsx"
    workbook = None
    try:
        generate_penalty_report([rec], str(output_path))
        workbook = openpyxl.load_workbook(output_path, read_only=False)
        ws = workbook["Comparison"]

        assert ws.cell(row=1, column=1).value == "Reissue and Refund Charges"
        assert ws.cell(row=2, column=4).value == "Re-issue"
        assert ws.cell(row=2, column=8).value == "Refund"
        assert ws.cell(row=3, column=4).value == "Before 24 Hours"
        assert ws.cell(row=3, column=10).value == "NOSHOW"
        assert ws.cell(row=4, column=1).value == "DAC-CCU (USD)"
        assert ws.cell(row=4, column=2).value == "BS"
        assert ws.cell(row=4, column=3).value == "ALL"
        assert ws.cell(row=4, column=4).value == "USD 10"
        assert ws.cell(row=4, column=6).value == "USD 30"
        assert ws.cell(row=4, column=8).value == "USD 20"
        assert ws.cell(row=4, column=10).value == "USD 50"
    finally:
        if workbook is not None:
            workbook.close()
            workbook = None
            gc.collect()
        if output_path.exists():
            output_path.unlink()

