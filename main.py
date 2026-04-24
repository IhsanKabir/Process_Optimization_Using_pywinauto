"""
main.py - Travelport Fare Automation Orchestrator

Phase 1 (MVP): Read raw GDS output from text files, parse, and generate Excel report.

Usage:
    python main.py                     # Process all raw files in data/raw/
    python main.py --no-changes        # Skip change detection
    python main.py --output report.xlsx # Custom output path
"""

# Must be first — before any UI or automation imports — so physical pixel
# coordinates are used consistently by both pyautogui and UIAutomation.
try:
    import ctypes as _ctypes
    _ctypes.windll.shcore.SetProcessDpiAwareness(2)
except Exception:
    pass

import argparse
import json
import logging
import os
import re
import shutil
import sys
import time as _time
import constants
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from datetime import datetime, timedelta
from collections import OrderedDict
from functools import lru_cache

try:
    import airportsdata
except ImportError:
    airportsdata = None

try:
    from tqdm import tqdm
except ImportError:
    # Fallback if tqdm is not installed
    tqdm = None

from parser import (
    parse_fare_display,
    group_fares_by_rbd,
    load_commands,
    generate_file_key,
    parse_command,
)
from penalty_parser import parse_penalty_text
from penalty_report import generate_penalty_report
from tax_breakdown_parser import parse_fs_tax_breakdown, looks_like_fs_tax_breakdown
from excel_report import generate_report
from change_detector import (
    detect_changes,
    detect_tax_changes,
    format_tax_change_summary,
    save_snapshot,
    snapshot_has_changed,
    load_latest_snapshot,
    load_snapshot_by_reference,
    format_change_summary,
)
from exceptions import ConfigurationError, ValidationError
from validators import (
    validate_config,
    validate_airline_code,
    validate_airport_code,
    validate_limit,
    validate_route,
    sanitize_command,
    validate_parsed_fares,
    validate_currency_code,
)
from credential_manager import CredentialManager
from checkpoint_manager import CheckpointManager
from constants import (
    MAX_RETRIES_COMMAND,
    FS_DATE_OFFSET_START,
    FS_DATE_STEP,
    FS_DATE_WINDOW_DAYS,
    FS_DATE_FALLBACK_OFFSET,
)

try:
    from database import DatabaseManager
except ImportError:
    DatabaseManager = None

try:
    import bigquery_pusher as _bq_pusher
except ImportError:
    _bq_pusher = None

# Try to load .env file if python-dotenv is available
try:
    from dotenv import load_dotenv

    load_dotenv(override=True)
except ImportError:
    pass  # python-dotenv not installed, skip

logger = logging.getLogger("travelport")

if getattr(sys, "frozen", False):
    # PyInstaller creates a temp folder and stores path in _MEIPASS
    # But we want to output files next to the installed executable
    SCRIPT_DIR = os.path.dirname(sys.executable)
else:
    SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

DEFAULT_CONFIG = os.path.join(SCRIPT_DIR, "config.json")
RAW_DATA_DIR = os.path.join(SCRIPT_DIR, "data", "raw")
REPORTS_DIR = os.path.join(SCRIPT_DIR, "data", "reports")
ARCHIVE_DIR = os.path.join(SCRIPT_DIR, "data", "archive")
LOG_DIR = os.path.join(SCRIPT_DIR, "data", "logs")
CHECKPOINT_DIR = os.path.join(SCRIPT_DIR, "data", "checkpoints")
RAW_PENALTY_DIR = os.path.join(SCRIPT_DIR, "data", "raw_penalty")
PLACEHOLDER_DATABASE_URLS = {
    "postgresql://user:password@localhost/travelport_db",
    "postgresql://user:password@localhost/GDS_Automation",
}
# Remote sources - admin updates these files on GitHub; all users get the
# latest config automatically on next run without needing a new exe.
REMOTE_CONFIG_URL = (
    "https://raw.githubusercontent.com/IhsanKabir/"
    "Process_Optimization_Using_pywinauto/main/config.json"
)
REMOTE_COMMANDS_URL = (
    "https://raw.githubusercontent.com/IhsanKabir/"
    "Process_Optimization_Using_pywinauto/main/commands.txt"
)
DEFAULT_COMMANDS_TEMPLATE = """# TravelportAuto route commands
# Add one Fare Display command per line.
# Examples:
# FDDACMCT/BG
# FDDACBKK/BS
# FDCGPMLE/BG
"""
FS_OPTION_PATTERN = re.compile(
    r"PRICING\s+OPTION\s+(\d+)(.*?(?=PRICING\s+OPTION\s+\d+|$))",
    re.IGNORECASE | re.DOTALL,
)
FS_LEG_PATTERN = re.compile(r"^\s*(\d+)\s+[#@-]?([A-Z0-9]{2})\s+", re.MULTILINE)
_RE_SAFE_FILENAME = re.compile(r"[^A-Za-z0-9._-]+")
DEFAULT_AIRPORT_COUNTRY_CODES = {
    "DAC": "BD",
    "CGP": "BD",
    "ZYL": "BD",
    "CXB": "BD",
    "MLE": "MV",
    "CAN": "CN",
    "MCT": "OM",
    "DOH": "QA",
    "DXB": "AE",
    "AUH": "AE",
    "SHJ": "AE",
    "RKT": "AE",
    "RUH": "SA",
    "JED": "SA",
    "DMM": "SA",
    "MED": "SA",
    "KWI": "KW",
    "BOM": "IN",
    "DEL": "IN",
    "MAA": "IN",
    "BLR": "IN",
    "CCU": "IN",
    "SIN": "SG",
    "BKK": "TH",
    "KUL": "MY",
    "HKT": "TH",
    "MNL": "PH",
    "SGN": "VN",
    "HAN": "VN",
    "PEK": "CN",
    "PVG": "CN",
    "SZX": "CN",
}


def setup_logging():
    """Configure logging to console (INFO) and file (DEBUG)."""
    os.makedirs(LOG_DIR, exist_ok=True)
    log_file = os.path.join(
        LOG_DIR, f"run_{datetime.now().strftime('%Y-%m-%d_%H%M')}.log"
    )

    root = logging.getLogger("travelport")
    root.setLevel(logging.DEBUG)

    # Console handler - skip when stdout has been replaced (e.g. by the GUI),
    # otherwise messages would be delivered twice: once via this StreamHandler
    # (writing to the redirected stdout -> queue) and once via the QueueHandler
    # already attached to this logger by the GUI.
    if sys.stdout is sys.__stdout__:
        ch = logging.StreamHandler()
        ch.setLevel(logging.INFO)
        ch.setFormatter(logging.Formatter("%(message)s"))
        root.addHandler(ch)

    # File: DEBUG level, full format with timestamps
    fh = logging.FileHandler(log_file, encoding="utf-8")
    fh.setLevel(logging.DEBUG)
    fh.setFormatter(
        logging.Formatter("%(asctime)s [%(levelname)s] %(name)s: %(message)s")
    )
    root.addHandler(fh)

    return log_file


def _resolve_bundled_file(filename: str) -> str | None:
    """Return a bundled file path when running from a PyInstaller build."""
    bundle_dir = getattr(sys, "_MEIPASS", None)
    if not bundle_dir:
        return None

    candidate = os.path.join(bundle_dir, filename)
    if os.path.exists(candidate):
        return candidate
    return None


def _seed_runtime_config_if_missing(config_path: str) -> str:
    """
    Seed the runtime config from the bundled default config when available.

    This keeps the packaged app usable even if a desktop only receives the exe.
    """
    if os.path.exists(config_path):
        return config_path

    bundled_config = _resolve_bundled_file("config.json")
    if not bundled_config:
        return config_path

    try:
        os.makedirs(os.path.dirname(config_path), exist_ok=True)
        shutil.copyfile(bundled_config, config_path)
        logger.info(f"  Created default config.json at {config_path}")
        return config_path
    except OSError as e:
        logger.warning(
            f"  Could not create local config.json ({e}). Using bundled defaults."
        )
        return bundled_config


def _fetch_remote_config() -> dict | None:
    """Fetch the latest config.json from GitHub. Returns None if offline or invalid."""
    import urllib.request as _ur

    try:
        req = _ur.Request(
            REMOTE_CONFIG_URL, headers={"User-Agent": "TravelportAuto/1.0"}
        )
        with _ur.urlopen(req, timeout=8) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except Exception:
        return None


def _ensure_commands_template(commands_file: str) -> bool:
    """Create a starter commands.txt file for first-run users if missing."""
    if os.path.exists(commands_file):
        return False

    try:
        parent = os.path.dirname(commands_file)
        if parent:
            os.makedirs(parent, exist_ok=True)
        with open(commands_file, "w", encoding="utf-8") as f:
            f.write(DEFAULT_COMMANDS_TEMPLATE)
        logger.info(f"  Created starter commands file at {commands_file}")
        return True
    except OSError as e:
        logger.warning(f"  Could not create starter commands file: {e}")
        return False


def _has_writable_stream(stream) -> bool:
    """Return True if tqdm can safely write to this stream."""
    return stream is not None and callable(getattr(stream, "write", None))


def _should_use_tqdm(gui_mode: bool = False) -> bool:
    """Enable tqdm only when a usable console-like stream exists."""
    if tqdm is None or gui_mode:
        return False
    return _has_writable_stream(sys.stderr) or _has_writable_stream(sys.stdout)


def _tqdm_stream():
    """Choose the best available stream for tqdm output."""
    if _has_writable_stream(sys.stderr):
        return sys.stderr
    if _has_writable_stream(sys.stdout):
        return sys.stdout
    return None


def _format_eta(seconds: float) -> str:
    """Format elapsed/remaining seconds as `Xm Ys` (or `Ys` when under a minute)."""
    seconds = max(0, int(round(seconds)))
    if seconds < 60:
        return f"{seconds}s"
    minutes, secs = divmod(seconds, 60)
    if minutes < 60:
        return f"{minutes}m {secs}s"
    hours, minutes = divmod(minutes, 60)
    return f"{hours}h {minutes}m {secs}s"


def _log_eta(
    logger_obj,
    start_ts: float,
    done: int,
    total: int,
    label: str = "progress",
    every: int = 5,
) -> None:
    """Emit `[ETA ~Xm Ys] label done/total` every `every` items (CLI fallback for tqdm).

    Skip if nothing done yet, total unknown, or we're not on a tick boundary.
    """
    import time as _t

    if total <= 0 or done <= 0:
        return
    if done != total and done % max(1, every) != 0:
        return
    elapsed = _t.time() - start_ts
    if elapsed <= 0:
        return
    per_item = elapsed / done
    remaining = per_item * (total - done)
    logger_obj.info(
        f"  [ETA ~{_format_eta(remaining)}] {label} {done}/{total} "
        f"(elapsed {_format_eta(elapsed)})"
    )


def _install_cli_esc_listener(stop_event, logger_obj) -> None:
    """Windows-only: poll GetAsyncKeyState(VK_ESCAPE) and set stop_event on press.

    No-op when stop_event is None, on non-Windows platforms, or when a GUI
    is driving the run (the GUI already owns the ESC shortcut).
    """
    if stop_event is None or sys.platform != "win32":
        return
    try:
        import ctypes
        import threading

        user32 = ctypes.windll.user32
        VK_ESCAPE = 0x1B

        def _watch() -> None:
            import time as _t

            while not stop_event.is_set():
                try:
                    if user32.GetAsyncKeyState(VK_ESCAPE) & 0x8000:
                        logger_obj.info(
                            "  [ESC] Cancel requested - finishing current step..."
                        )
                        stop_event.set()
                        return
                except Exception:
                    return
                _t.sleep(0.1)

        t = threading.Thread(target=_watch, daemon=True, name="cli-esc-listener")
        t.start()
    except Exception as exc:  # pragma: no cover - best effort
        logger_obj.debug(f"  CLI ESC listener unavailable: {exc}")


def _env_truthy(name: str) -> bool:
    """Return True when an environment variable is set to a truthy value."""
    return str(os.getenv(name, "")).strip().lower() in {"1", "true", "yes", "on"}


def load_config(config_path: str) -> dict:
    """Load config - remote GitHub first (auto-updates), local/bundled fallback."""
    # Try remote first - this keeps airline names, airport lists, etc. current
    # without requiring users to update the exe or edit any files.
    remote = _fetch_remote_config()
    if remote is not None:
        try:
            config = validate_config(remote)
            logger.debug("  Config loaded from remote (ok)")
            return config
        except ConfigurationError:
            logger.warning("  Remote config failed validation - falling back to local.")

    # Fall back to bundled / local config.json
    if os.path.abspath(config_path) == os.path.abspath(DEFAULT_CONFIG):
        config_path = _seed_runtime_config_if_missing(config_path)

    try:
        with open(config_path, "r", encoding="utf-8") as f:
            config = json.load(f)
    except FileNotFoundError:
        logger.error(f"  Config file not found: {config_path}")
        raise ConfigurationError(f"Config file not found: {config_path}")
    except json.JSONDecodeError as e:
        logger.error(f"  Invalid JSON in config file: {e}")
        raise ConfigurationError(f"Invalid JSON in config file: {e}")

    # Validate configuration using validators module
    try:
        config = validate_config(config)
        logger.debug("  Configuration validated successfully")
    except ConfigurationError as e:
        logger.error(f"  Configuration validation failed: {e}")
        raise

    return config


def load_raw_data(raw_dir: str) -> tuple[dict[str, str], dict[str, str]]:
    """Load raw GDS output from text files in data/raw/."""
    raw_texts = {}
    raw_fs_texts = {}

    if not os.path.exists(raw_dir):
        return raw_texts, raw_fs_texts

    for filename in sorted(os.listdir(raw_dir)):
        if (
            filename.endswith(".txt")
            and not filename.startswith("_")
            and not filename.startswith("FTAX")
        ):
            filepath = os.path.join(raw_dir, filename)
            with open(filepath, "r", encoding="utf-8") as f:
                content = f.read()

            if filename.endswith("_FS.txt"):
                file_key = filename[:-7]
                raw_fs_texts[file_key] = content
            else:
                file_key = filename[:-4]
                raw_texts[file_key] = content

            logger.info(f"  Loaded: {filename}")

    return raw_texts, raw_fs_texts


def process_route_data(
    raw_texts: dict[str, str],
    raw_fs_texts: dict[str, str],
    config: dict,
    enable_validation: bool = True,
    show_progress: bool = True,
    stop_event=None,
) -> dict:
    """Process raw text files into route data with optional validation."""
    rbd_sort_order = config.get("rbd_sort_order", [])
    all_route_data = OrderedDict()
    progress_active = show_progress and _should_use_tqdm()

    # Combine all file_keys from both raw_texts and raw_fs_texts
    all_file_keys = set(raw_texts.keys()) | set(raw_fs_texts.keys())

    # Use tqdm progress bar if available
    items = sorted(all_file_keys)
    if progress_active:
        items = tqdm(items, desc="Processing routes", unit="route", file=_tqdm_stream())

    for file_key in items:
        if stop_event and stop_event.is_set():
            logger.info("  [STOP] Stop requested during parsing. Stopping early.")
            break
        raw_text = raw_texts.get(file_key, "")

        # Parse fare data if available
        fares = []
        currency = None
        if raw_text:
            result = parse_fare_display(raw_text)
            fares = result.get("fares", [])
            currency = result.get("currency")

            # Validate currency code
            if enable_validation and currency:
                validate_currency_code(currency, warn_only=True)

            # Validate parsed fares
            if enable_validation and fares:
                validation_stats = validate_parsed_fares(fares, currency)
                if validation_stats["invalid_fares"] > 0:
                    logger.warning(
                        f"  {file_key}: {validation_stats['invalid_fares']}/{validation_stats['total_fares']} "
                        f"fares have validation issues"
                    )

        # Parse FS tax data if available
        fs_taxes = {}
        if file_key in raw_fs_texts:
            fs_taxes = parse_fs_tax_breakdown(raw_fs_texts[file_key])

            # Debug logging for tax parsing
            if not progress_active:
                if fs_taxes.get("exchange_rate", 0) > 0:
                    logger.debug(
                        f"      [TAX] {file_key}: Rate={fs_taxes.get('exchange_rate'):.4f}, "
                        f"Total={fs_taxes.get('total_taxes', 0)}, YQ={fs_taxes.get('yq_charge', 0)}"
                    )
                else:
                    logger.warning(
                        f"      [TAX] {file_key}: Exchange rate is 0 or missing - tax data may not display correctly"
                    )
                    logger.warning(
                        f"      [TAX] Base fare={fs_taxes.get('base_fare', 0)}, Equ fare={fs_taxes.get('equ_fare', 0)}"
                    )
                    if raw_fs_texts[file_key]:
                        logger.warning(
                            f"      [TAX] First 200 chars of raw FS text: {raw_fs_texts[file_key][:200]}"
                        )

        # Add to all_route_data if we have either fares or taxes
        if fares or fs_taxes:
            grouped = group_fares_by_rbd(fares, rbd_sort_order) if fares else {}
            all_route_data[file_key] = {
                "rbd_data": grouped,
                "currency": currency,
                "fs_taxes": fs_taxes,
            }

            # Only log if not using tqdm (to avoid cluttering progress bar)
            if not progress_active:
                if fares:
                    ow_count = sum(
                        1 for d in grouped.values() if d.get("ow_fare") is not None
                    )
                    rt_count = sum(
                        1 for d in grouped.values() if d.get("rt_fare") is not None
                    )
                    logger.info(
                        f"  {file_key} -> {len(grouped)} RBDs ({ow_count} OW, {rt_count} RT) [{currency or 'N/A'}]"
                    )
                elif fs_taxes:
                    logger.info(f"  {file_key} -> Tax data only (no fares)")
        else:
            if not progress_active:
                logger.warning(f"  No data parsed from: {file_key}")

    return OrderedDict(sorted(all_route_data.items()))


def record_to_database(all_route_data: dict, config: dict, mode: str = "auto") -> int:
    """Persist the data to PostgreSQL if configured. Returns run_id (>0) or -1."""
    db_url = _resolve_database_url(config)
    if not db_url or "postgresql://" not in db_url:
        return -1

    if not DatabaseManager:
        logger.warning(
            "  [!] Database integration skipped (psycopg2-binary not installed)"
        )
        return -1

    logger.info(f"[DB] Recording {len(all_route_data)} items to history...")
    db = DatabaseManager(db_url)
    if db.connect():
        if mode == "tax-mode":
            run_id = db.record_tax_run(all_route_data, run_mode=mode)
        else:
            run_id = db.record_run(all_route_data, run_mode=mode)
        if run_id > 0:
            logger.info(f"  [OK] Database record created (Run ID: {run_id})")
        db.close()
        return run_id
    else:
        logger.warning("  [!] Database integration skipped (connection failed)")
        return -1


def show_usage():
    """Show usage instructions when no data files are found."""
    logger.info("")
    logger.info("  [!] No .txt files found in data/raw/")
    logger.info("")
    logger.info("  HOW TO USE:")
    logger.info("  " + "-" * 44)
    logger.info("  1. Open Travelport Smartpoint")
    logger.info("  2. Run a fare display command (e.g., FDDACMLE/BG)")
    logger.info("  3. Select all output and copy (Ctrl+C)")
    logger.info(f"  4. Save as .txt in: {RAW_DATA_DIR}")
    logger.info("  5. Run: python main.py")


def _process_airport_taxes(
    automation,
    airports: dict,
    failed_commands: list,
    stop_event=None,
    use_tqdm: bool = False,
    write_backups: bool = False,
    raw_data_dir: str = "",
):
    """Shared logic for extracting FTAX data for a set of airports.

    Returns dict: {airport_code: {"taxes": [detail_data, ...]}}
    """
    from tax_parser import parse_ftax_list, parse_ftax_detail

    result = {}
    airport_items = airports.items()
    if use_tqdm and tqdm:
        airport_items = tqdm(
            list(airport_items),
            desc="Extracting tax data",
            unit="airport",
            file=_tqdm_stream(),
        )
    else:
        airport_items = list(airport_items)

    for index, (airport_code, airport_info) in enumerate(airport_items, 1):
        if stop_event and stop_event.is_set():
            logger.info("  [STOP] Stop requested - finishing after this point.")
            break
        country_code = airport_info["country"]

        if not use_tqdm:
            logger.info(
                f"  [{index}/{len(airports)}] Airport: {airport_code} ({country_code})"
            )

        list_cmd = f"FTAX-{country_code}"
        list_text = automation.run_command(list_cmd, max_pages=1)

        if write_backups and raw_data_dir:
            os.makedirs(raw_data_dir, exist_ok=True)
            backup_path = os.path.join(raw_data_dir, f"{list_cmd}.txt")
            with open(backup_path, "w", encoding="utf-8") as f:
                f.write(list_text)

        tax_types = parse_ftax_list(list_text)
        logger.info(
            f"    Found {len(tax_types)} tax types: {', '.join(t['code'] for t in tax_types)}"
        )

        if not tax_types and write_backups:
            logger.warning(
                f"    [DEBUG] Raw list text ({len(list_text)} chars): {list_text[:300]}"
            )

        airport_tax_details = []
        for idx, t in enumerate(tax_types, 1):
            detail_text = automation.run_ftax_command(
                country_code, t["code"], tax_index=idx
            )

            if write_backups and raw_data_dir:
                detail_backup = os.path.join(
                    raw_data_dir, f"FTAX-{country_code}_{t['code']}.txt"
                )
                with open(detail_backup, "w", encoding="utf-8") as f:
                    f.write(detail_text)

            if not detail_text or len(detail_text.strip()) < 20:
                failed_commands.append(f"{country_code}/{t['code']}")
                logger.warning(f"    Failed to extract details for {t['code']}")
                continue

            if write_backups:
                logger.debug(
                    f"      [DEBUG] Raw detail text first 200 chars: {detail_text[:200]}"
                )

            detail_data = parse_ftax_detail(detail_text, t["code"], t["name"])
            airport_tax_details.append(detail_data)
            logger.info(
                f"      {t['code']} -> {len(detail_data['sections'])} sections extracted."
            )

            if not detail_data["sections"] and write_backups:
                logger.warning(
                    f"      [DEBUG] 0 sections! Full text ({len(detail_text)} chars): {detail_text[:500]}"
                )

            if idx < len(tax_types):
                automation.return_to_tax_list(country_code)

        result[airport_code] = {
            "taxes": airport_tax_details,
            "_country": country_code,
        }

    return result


def _safe_filename(value: str) -> str:
    """Convert arbitrary text into a filesystem-safe filename fragment."""
    cleaned = _RE_SAFE_FILENAME.sub("_", str(value or ""))
    return cleaned.strip("_") or "unknown"


def _is_placeholder_database_url(database_url: str | None) -> bool:
    normalized = (database_url or "").strip()
    if not normalized:
        return True
    if normalized in PLACEHOLDER_DATABASE_URLS:
        return True
    return "YOUR_" in normalized.upper()


def _normalize_database_url(database_url: str | None) -> str:
    """Normalize PostgreSQL URLs so psycopg2 can consume them reliably."""
    normalized = (database_url or "").strip()
    if not normalized:
        return ""

    if normalized.startswith("postgresql+"):
        normalized = "postgresql://" + normalized.split("://", 1)[1]

    if normalized.startswith("postgresql://"):
        parts = urlsplit(normalized)
        query = dict(parse_qsl(parts.query, keep_blank_values=True))
        query.setdefault("connect_timeout", "5")
        normalized = urlunsplit(
            (
                parts.scheme,
                parts.netloc,
                parts.path,
                urlencode(query),
                parts.fragment,
            )
        )

    return normalized


def _resolve_database_url(config: dict) -> str:
    """Prefer DATABASE_URL, but fall back to config when it isn't a placeholder."""
    env_database_url = _normalize_database_url(os.environ.get("DATABASE_URL"))
    if env_database_url:
        return env_database_url

    config_database_url = config.get("database_url")
    if _is_placeholder_database_url(config_database_url):
        return ""

    return _normalize_database_url(config_database_url)


def _route_variants(route_code: str, one_direction: bool = False) -> set[str]:
    """Return the route directions that should be treated as a match."""
    normalized = (route_code or "").strip().upper().replace("-", "")
    if len(normalized) != 6:
        return set()
    if one_direction:
        return {normalized}
    return {normalized, f"{normalized[3:]}{normalized[:3]}"}


def _command_matches_route(
    command_entry: dict, route_code: str, one_direction: bool = False
) -> bool:
    """Check whether a command matches a route filter."""
    route_variants = _route_variants(route_code, one_direction=one_direction)
    if not route_variants:
        return False

    origin = str(command_entry.get("origin") or "").upper()
    destination = str(command_entry.get("destination") or "").upper()
    if origin and destination:
        return f"{origin}{destination}" in route_variants

    raw_command = str(command_entry.get("command") or "").upper().replace("-", "")
    return any(route_variant in raw_command for route_variant in route_variants)


def _parse_requested_routes(route_query: str | None) -> list[str]:
    """Normalize comma-separated route input into canonical `AAA-BBB` strings."""
    if not route_query:
        return []

    normalized_routes: list[str] = []
    seen_routes: set[str] = set()
    for raw_route in str(route_query).split(","):
        raw_route = raw_route.strip()
        if not raw_route:
            continue
        origin, destination = validate_route(raw_route)
        route_code = f"{origin}-{destination}"
        if route_code in seen_routes:
            continue
        seen_routes.add(route_code)
        normalized_routes.append(route_code)
    return normalized_routes


def _resolve_explicit_airline_codes(
    airline_query: str | None, config: dict
) -> list[str]:
    """Resolve explicit fare/penalty airline filters or fall back to configured airlines."""
    if airline_query:
        airline_codes: list[str] = []
        seen_codes: set[str] = set()
        for raw_code in str(airline_query).split(","):
            raw_code = raw_code.strip()
            if not raw_code:
                continue
            code = validate_airline_code(raw_code)
            if code in seen_codes:
                continue
            seen_codes.add(code)
            airline_codes.append(code)
        if airline_codes:
            return airline_codes

    airline_names = config.get("airline_names", {})
    if isinstance(airline_names, dict) and airline_names:
        return [str(code).upper() for code in airline_names.keys()]

    raise ConfigurationError(
        "No airlines available to generate explicit route commands. Add airline_names or provide --airline."
    )


def _build_explicit_route_commands(
    route_query: str,
    airline_query: str | None,
    config: dict,
    one_direction: bool = False,
) -> tuple[list[dict], list[str], list[str]]:
    """Build FD commands directly from typed routes and airlines."""
    routes = _parse_requested_routes(route_query)
    if not routes:
        return [], [], []

    airline_codes = _resolve_explicit_airline_codes(airline_query, config)
    commands: list[dict] = []
    seen_commands: set[str] = set()

    for route_code in routes:
        origin, destination = route_code.split("-", 1)
        direction_pairs = [(origin, destination)]
        if not one_direction:
            direction_pairs.append((destination, origin))

        for dir_origin, dir_destination in direction_pairs:
            for airline_code in airline_codes:
                command = f"FD{dir_origin}{dir_destination}/{airline_code}"
                if command in seen_commands:
                    continue
                seen_commands.add(command)
                commands.append(
                    {
                        "origin": dir_origin,
                        "destination": dir_destination,
                        "airline": airline_code,
                        "route": f"{dir_origin}-{dir_destination}",
                        "command": command,
                    }
                )

    return commands, routes, airline_codes


def _extract_tax_airport_queries(
    airport_query: str | None = None, route_query: str | None = None
) -> list[str] | None:
    """Normalize explicit tax-airport input from `--airport` or legacy `--route`."""
    raw_query = (airport_query or "").strip() or (route_query or "").strip()
    if not raw_query:
        return None

    normalized_queries: list[str] = []
    for chunk in raw_query.split(","):
        chunk = chunk.strip()
        if not chunk:
            continue

        if "-" in chunk:
            route_candidate = re.sub(r"\s+", "", chunk)
            try:
                origin, _destination = validate_route(route_candidate)
                normalized_queries.append(origin)
                continue
            except ValidationError:
                pass

        normalized_queries.append(chunk)

    return normalized_queries or None


def _tax_airport_display_name(airport_code: str, airport_info: dict, config: dict) -> str:
    """Return the most user-friendly label for a configured tax airport."""
    city_name = str(config.get("city_names", {}).get(airport_code, "") or "").strip()
    if city_name:
        return city_name
    city = str(airport_info.get("city", "") or "").strip()
    if city:
        return city
    info_name = str(airport_info.get("name", "") or "").strip()
    if info_name:
        return info_name
    return airport_code


def _tax_airport_name_aliases(
    airport_code: str, airport_info: dict, config: dict
) -> list[str]:
    """Return configured non-code aliases that can identify a tax airport."""
    aliases: list[str] = []
    for alias in (
        config.get("city_names", {}).get(airport_code),
        airport_info.get("city"),
        airport_info.get("name"),
    ):
        text = str(alias or "").strip()
        if not text:
            continue
        if text.casefold() == airport_code.casefold():
            continue
        if any(text.casefold() == existing.casefold() for existing in aliases):
            continue
        aliases.append(text)
    return aliases


@lru_cache(maxsize=1)
def _load_global_airport_directory() -> dict[str, dict]:
    """Load a global IATA airport directory for explicit Future Tax searches."""
    if airportsdata is None:
        return {}

    searchable: dict[str, dict] = {}
    for airport_code, airport_info in airportsdata.load("IATA").items():
        normalized_code = str(airport_code or "").upper().strip()
        try:
            normalized_code = validate_airport_code(normalized_code)
        except ValidationError:
            continue

        country_code = str(airport_info.get("country", "") or "").upper().strip()
        if not country_code:
            continue

        entry = {
            "country": country_code,
            "_source": "global",
        }
        city_name = str(airport_info.get("city", "") or "").strip()
        if city_name:
            entry["city"] = city_name
        airport_name = str(airport_info.get("name", "") or "").strip()
        if airport_name:
            entry["name"] = airport_name

        searchable[normalized_code] = entry

    return searchable


def _configured_airport_country_codes(config: dict) -> dict[str, str]:
    """Return known airport -> country code mappings for explicit tax searches."""
    country_codes = dict(DEFAULT_AIRPORT_COUNTRY_CODES)
    for airport_code, country_code in config.get("airport_country_codes", {}).items():
        country_codes[str(airport_code).upper()] = str(country_code).upper()
    for airport_code, airport_info in config.get("tax_airports", {}).items():
        country_code = str(airport_info.get("country", "") or "").upper()
        if country_code:
            country_codes[str(airport_code).upper()] = country_code
    return country_codes


def _build_searchable_tax_airports(config: dict) -> dict[str, dict]:
    """Return airports that can be used for explicit Future Tax lookups."""
    searchable = {
        airport_code: dict(airport_info)
        for airport_code, airport_info in _load_global_airport_directory().items()
    }

    for airport_code, country_code in _configured_airport_country_codes(config).items():
        existing = dict(searchable.get(airport_code, {}))
        existing["country"] = country_code
        if not existing.get("_source"):
            existing["_source"] = "config"
        searchable[airport_code] = existing

    for airport_code, city_name in config.get("city_names", {}).items():
        normalized_code = str(airport_code).upper()
        if normalized_code not in searchable:
            continue
        text = str(city_name or "").strip()
        if text:
            searchable[normalized_code]["city"] = text

    for airport_code, airport_info in config.get("tax_airports", {}).items():
        existing = dict(searchable.get(airport_code, {}))
        existing.update(dict(airport_info))
        existing["_source"] = "config"
        searchable[airport_code] = existing

    return searchable


def _format_tax_airport_candidates(
    airport_codes: list[str], tax_airports: dict, config: dict, max_candidates: int = 12
) -> str:
    """Return a compact list of configured tax-airport choices."""
    labels = []
    shown_codes = airport_codes[:max_candidates]
    for airport_code in shown_codes:
        airport_info = tax_airports.get(airport_code, {})
        display_name = _tax_airport_display_name(airport_code, airport_info, config)
        extra_name = str(airport_info.get("name", "") or "").strip()
        if extra_name and extra_name.casefold() != display_name.casefold():
            labels.append(f"{airport_code} ({display_name} / {extra_name})")
        else:
            labels.append(f"{airport_code} ({display_name})")
    if len(airport_codes) > max_candidates:
        labels.append(f"... ({len(airport_codes) - max_candidates} more)")
    return ", ".join(labels)


def _has_global_tax_airport_search(tax_airports: dict) -> bool:
    """Return True when the explicit tax resolver includes the global airport directory."""
    return any(
        str(airport_info.get("_source", "")).casefold() == "global"
        for airport_info in tax_airports.values()
    )


def _resolve_tax_airport_query(
    query: str, tax_airports: dict, config: dict
) -> tuple[str, dict, str]:
    """Resolve a single airport code or configured airport name to one tax airport."""
    query_text = str(query or "").strip()
    if not query_text:
        raise ValidationError("airport", query, "cannot be empty")

    query_key = query_text.casefold()
    search_scope = (
        "airports" if _has_global_tax_airport_search(tax_airports) else "configured tax airports"
    )

    for airport_code, airport_info in tax_airports.items():
        if airport_code.casefold() == query_key:
            return airport_code, airport_info, airport_code

    exact_name_matches: list[tuple[str, str]] = []
    partial_matches: list[tuple[str, str]] = []

    for airport_code, airport_info in tax_airports.items():
        aliases = _tax_airport_name_aliases(airport_code, airport_info, config)
        exact_alias = next(
            (alias for alias in aliases if alias.casefold() == query_key),
            None,
        )
        if exact_alias:
            exact_name_matches.append((airport_code, exact_alias))
            continue

        partial_alias = next(
            (alias for alias in aliases if query_key in alias.casefold()),
            None,
        )
        if partial_alias:
            partial_matches.append((airport_code, partial_alias))

    if len(exact_name_matches) == 1:
        airport_code, matched_alias = exact_name_matches[0]
        return airport_code, tax_airports[airport_code], matched_alias

    if len(exact_name_matches) > 1:
        matched_codes = [airport_code for airport_code, _alias in exact_name_matches]
        raise ValidationError(
            "airport",
            query,
            f"matches multiple {search_scope}: "
            + _format_tax_airport_candidates(matched_codes, tax_airports, config),
        )

    if len(partial_matches) == 1:
        airport_code, matched_alias = partial_matches[0]
        return airport_code, tax_airports[airport_code], matched_alias

    if len(partial_matches) > 1:
        matched_codes = [airport_code for airport_code, _alias in partial_matches]
        raise ValidationError(
            "airport",
            query,
            f"matches multiple {search_scope}: "
            + _format_tax_airport_candidates(matched_codes, tax_airports, config),
        )

    if _has_global_tax_airport_search(tax_airports):
        raise ValidationError(
            "airport",
            query,
            "not found in known airports. Try a 3-letter airport code like SYD or a more specific airport or city name.",
        )

    raise ValidationError(
        "airport",
        query,
        "not found in configured tax airports. Valid options: "
        + _format_tax_airport_candidates(sorted(tax_airports.keys()), tax_airports, config),
    )


def _configured_route_airports(commands_file: str) -> set[str]:
    """Return airport codes referenced by configured FD commands."""
    route_airports: set[str] = set()
    if not os.path.exists(commands_file):
        return route_airports

    with open(commands_file, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line.startswith("FD") and "/" in line and len(line) >= 8:
                route_airports.add(line[2:5].upper())
                route_airports.add(line[5:8].upper())
    return route_airports


def _filter_tax_airports_by_configured_routes(
    tax_airports: dict, commands_file: str
) -> tuple[dict, int]:
    """Filter tax airports to those appearing in configured routes when possible."""
    route_airports = _configured_route_airports(commands_file)
    if not route_airports:
        return tax_airports, 0

    filtered = {
        airport_code: airport_info
        for airport_code, airport_info in tax_airports.items()
        if airport_code.upper() in route_airports
    }
    if not filtered:
        return tax_airports, 0

    return filtered, len(tax_airports) - len(filtered)


def _select_tax_airports_for_run(config: dict, args) -> tuple[dict, dict]:
    """Select which configured tax airports should run for the current invocation."""
    tax_airports = dict(config.get("tax_airports", {}))
    searchable_tax_airports = _build_searchable_tax_airports(config)
    metadata = {
        "requested_queries": [],
        "resolved_codes": [],
        "resolutions": [],
        "skipped_by_route_filter": 0,
    }

    requested_queries = _extract_tax_airport_queries(
        getattr(args, "airport", None), getattr(args, "route", None)
    )
    if requested_queries:
        selected_tax_airports: dict = {}
        resolutions: list[dict] = []

        for requested_query in requested_queries:
            airport_code, airport_info, matched_alias = _resolve_tax_airport_query(
                requested_query, searchable_tax_airports, config
            )
            if airport_code in selected_tax_airports:
                continue

            selected_info = dict(airport_info)
            display_name = _tax_airport_display_name(airport_code, selected_info, config)
            selected_info["_display_name"] = display_name
            selected_info["_matched_alias"] = matched_alias
            selected_tax_airports[airport_code] = selected_info
            resolutions.append(
                {
                    "requested_query": requested_query,
                    "resolved_code": airport_code,
                    "display_name": display_name,
                    "matched_alias": matched_alias,
                    "country_code": selected_info.get("country"),
                }
            )

        metadata.update(
            {
                "requested_queries": requested_queries,
                "resolved_codes": list(selected_tax_airports.keys()),
                "resolutions": resolutions,
            }
        )
        return selected_tax_airports, metadata

    if not tax_airports:
        return tax_airports, metadata

    commands_file = os.path.join(SCRIPT_DIR, config.get("commands_file", "commands.txt"))
    tax_airports, skipped = _filter_tax_airports_by_configured_routes(
        tax_airports, commands_file
    )
    metadata["skipped_by_route_filter"] = skipped

    if getattr(args, "limit", 0) > 0:
        tax_airports = {
            airport_code: airport_info
            for index, (airport_code, airport_info) in enumerate(tax_airports.items())
            if index < args.limit
        }

    return tax_airports, metadata


def _fd_output_has_fares(raw_text: str) -> bool:
    """Return True when FD output contains actual fare rows."""
    if not raw_text or not raw_text.strip():
        return False
    parsed = parse_fare_display(raw_text)
    return bool(parsed.get("fares"))


def _fd_output_is_no_fares(raw_text: str) -> bool:
    """Return True when Smartpoint definitively says the FD request has no fares."""
    upper = (raw_text or "").upper()
    return "NO FARES FOUND" in upper and "INPUT REQUEST" in upper


def _fs_output_is_no_results(raw_text: str) -> bool:
    """Return True when FS checkout has returned a final no-result/error screen."""
    upper = (raw_text or "").upper()
    return any(
        marker in upper
        for marker in (
            "NO FARES FOUND",
            "CHECK ACTION CODE",
            "INVALID",
        )
    )


def _should_run_fs_extraction(args, fd_terminal_text: str) -> bool:
    """Skip FS when normal FD extraction produced no fare rows."""
    if getattr(args, "only_fd", False):
        return False
    if getattr(args, "only_yq", False) or getattr(args, "only_currency", False):
        return True
    return _fd_output_has_fares(fd_terminal_text)


def _find_pure_airline_option_in_fs_page(
    fs_text: str, airline: str
) -> tuple[int | None, str | None, int]:
    """Find the first pricing option on the current FS page containing only the target airline."""
    options = list(FS_OPTION_PATTERN.finditer(fs_text or ""))
    airline_upper = (airline or "").upper()

    for option_index, opt_match in enumerate(options):
        opt_num = opt_match.group(1)
        block = opt_match.group(2)
        leg_matches = FS_LEG_PATTERN.findall(block)
        if not leg_matches:
            continue

        leg_airlines = [match[1].upper().strip() for match in leg_matches]
        if all(code == airline_upper for code in leg_airlines):
            return option_index, opt_num, len(options)

    return None, None, len(options)


def _should_recheck_same_fs_page(
    current_fs_page: str, refreshed_fs_page: str, airline: str
) -> bool:
    """Return True when a same-date FS page looks more complete after settling."""
    current_text = (current_fs_page or "").strip()
    refreshed_text = (refreshed_fs_page or "").strip()
    if not refreshed_text or refreshed_text == current_text:
        return False

    (
        current_option_index,
        _current_option_number,
        current_option_count,
    ) = _find_pure_airline_option_in_fs_page(current_text, airline)
    (
        refreshed_option_index,
        _refreshed_option_number,
        refreshed_option_count,
    ) = _find_pure_airline_option_in_fs_page(refreshed_text, airline)

    if refreshed_option_index is not None and current_option_index is None:
        return True

    if refreshed_option_count > current_option_count:
        return True

    return (
        "PRICING OPTION" in refreshed_text.upper()
        and len(refreshed_text) > len(current_text) + 40
    )


def run_with_args(args, stop_event=None):
    """Entry point for the GUI: run extraction with a pre-built args Namespace."""
    if args.config is None:
        args.config = DEFAULT_CONFIG
    args._gui_mode = True
    return main(prebuilt_args=args, stop_event=stop_event)


def _run_currency_report_mode(args, config, stop_event):
    """Run standalone Currency Rate report via FZS for the configured pair list."""
    from datetime import date as _date, datetime as _dt

    from currency_archive import import_previous_rates
    from currency_report import generate_currency_report

    pairs = config.get("currency_report_pairs") or []
    local_cur = config.get("local_currency", "BDT")
    if not pairs:
        logger.error("  No 'currency_report_pairs' configured.")
        sys.exit(1)

    # Optional: seed a previous-day snapshot from a user file before generating.
    if getattr(args, "load_previous_rates", None):
        prev_date_str = getattr(args, "previous_date", None)
        if prev_date_str:
            try:
                prev_date = _dt.strptime(prev_date_str, "%Y-%m-%d").date()
            except ValueError:
                logger.error(f"  Invalid --previous-date value: {prev_date_str}")
                sys.exit(1)
        else:
            from datetime import timedelta
            prev_date = _date.today() - timedelta(days=1)
        try:
            snap = import_previous_rates(args.load_previous_rates, prev_date)
            logger.info(
                f"  [IMPORT] Loaded {len(snap.rates)} rates for {prev_date} from "
                f"{os.path.basename(args.load_previous_rates)}"
            )
        except (FileNotFoundError, ValueError, RuntimeError) as exc:
            logger.error(f"  Could not import previous rates: {exc}")
            sys.exit(1)

    if not args.auto:
        logger.error("  --currency-report requires --auto (live Smartpoint session).")
        sys.exit(1)

    logger.info(
        f"[2/3] Extracting FZS rates for {len(pairs)} currencies into {local_cur}..."
    )
    from smartpoint_automation import SmartpointAutomation, StopRequested
    from fzs_parser import parse_fzs_output

    automation = SmartpointAutomation(stop_event=stop_event)
    if not automation.connect():
        logger.error("  Please ensure Smartpoint is open.")
        sys.exit(1)

    username, password, pcc = CredentialManager.get_credentials()
    if username and password and not _env_truthy("SMARTPOINT_SKIP_AUTO_LOGIN"):
        try:
            if not automation.login(username, password, pcc):
                logger.error("  Login failed; sign in manually and retry.")
                sys.exit(1)
        except Exception as exc:
            logger.error(f"  Login error: {exc}")
            sys.exit(1)

    automation.refresh_terminal()

    current_rates: dict[str, float] = {}
    import time as _t_fzs

    fzs_start = _t_fzs.time()
    fzs_total = len(pairs)
    for idx, cur in enumerate(pairs, 1):
        if stop_event and stop_event.is_set():
            logger.info("  [STOP] Stop requested — finishing after this point.")
            break
        cur = cur.upper()
        try:
            raw = automation.run_fzs_command(cur, local_cur)
        except StopRequested:
            logger.info("  [STOP] Stop requested during FZS extraction.")
            break
        except Exception as exc:
            logger.warning(f"  [FZS] {cur}->{local_cur} failed: {exc}")
            continue
        parsed = parse_fzs_output(raw, cur, local_cur)
        if parsed.get("rate"):
            current_rates[cur] = float(parsed["rate"])
            logger.info(f"  [{idx}/{fzs_total}] {cur} -> {local_cur}: {parsed['rate']}")
        else:
            logger.warning(f"  [{idx}/{fzs_total}] Could not parse rate for {cur}")
        _log_eta(logger, fzs_start, idx, fzs_total, label="FZS rates", every=5)

    if not current_rates:
        logger.error("  No rates extracted; report not generated.")
        sys.exit(1)

    os.makedirs(REPORTS_DIR, exist_ok=True)
    ts = _dt.now().strftime("%Y-%m-%d_%H%M")
    if args.output:
        out_path = args.output
    else:
        out_path = os.path.join(REPORTS_DIR, f"currency_rates_{ts}.xlsx")

    logger.info("[3/3] Writing Currency Rate report...")
    result = generate_currency_report(current_rates, out_path)
    logger.info(f"  [OK] Currency Rate report: {result.path}")
    logger.info(
        f"  Previous-day basis: "
        f"{result.previous_date.isoformat() if result.previous_date else 'none (first run)'}"
    )
    try:
        from usage_tracker import increment as _inc
        _inc("currency_runs")
    except Exception:
        pass
    return result.path


def main(prebuilt_args=None, stop_event=None):
    """Main entry point. Called directly from CLI or via run_with_args() from GUI."""
    import time as _time

    start_time = _time.time()
    _stop = stop_event  # shorthand used in command loops below

    arg_parser = argparse.ArgumentParser(description="Travelport Data Automation")
    arg_parser.add_argument("--config", "-c", default=DEFAULT_CONFIG)
    arg_parser.add_argument(
        "--no-changes", action="store_true", help="Skip change detection"
    )
    arg_parser.add_argument("--output", "-o", default=None, help="Output Excel path")
    arg_parser.add_argument(
        "--auto", action="store_true", help="Extract data from live Smartpoint"
    )
    arg_parser.add_argument(
        "--limit", type=int, default=0, help="Limit commands (testing)"
    )
    arg_parser.add_argument(
        "--route",
        type=str,
        help="Filter fares to a specific route (e.g. DAC-MLE). In --tax mode, this remains a backward-compatible single-airport alias.",
    )
    arg_parser.add_argument(
        "--airport",
        type=str,
        help="In --tax mode, run one or more airport codes or airport names, comma-separated (e.g. KUL,SYD or Kuala Lumpur,Muscat). Blank keeps the configured tax-airport list.",
    )
    arg_parser.add_argument(
        "-1d",
        "--one-direction",
        action="store_true",
        help="With --route, match only the exact route direction instead of both directions",
    )
    arg_parser.add_argument(
        "--airline",
        type=str,
        help="Filter to specific airline(s), comma-separated (e.g. BG or BG,BS)",
    )
    arg_parser.add_argument(
        "--only-fd",
        action="store_true",
        help="Extract only basic Fares (skip YQ/Currency FS command)",
    )
    arg_parser.add_argument(
        "--only-yq",
        action="store_true",
        help="Extract only YQ and Tax Breakdown (skip Fares)",
    )
    arg_parser.add_argument(
        "--only-currency",
        action="store_true",
        help="Extract only exchange rates (alias for --only-yq)",
    )
    arg_parser.add_argument(
        "--tax",
        action="store_true",
        help="Extract Future Tax (FTAX) data instead of fares; optionally target one or more airports with --airport",
    )
    arg_parser.add_argument(
        "--penalty",
        action="store_true",
        help="Extract Rule 16 penalties per fare basis in a separate run",
    )
    arg_parser.add_argument(
        "--include-ftax",
        action="store_true",
        help="Extract global FTAX data alongside the specific route fares",
    )
    arg_parser.add_argument(
        "--quick-paste",
        action="store_true",
        help="Interactive mode: paste terminal texts physically from clipboard",
    )
    arg_parser.add_argument(
        "--speed",
        type=str,
        choices=["fast", "safe"],
        default=None,
        help='Speed profile: "fast" (aggressive timings, ~50%% faster) or "safe" (conservative timings for slower machines)',
    )
    arg_parser.add_argument(
        "--checkpoint",
        action="store_true",
        help="Enable checkpoint/resume mode for long runs",
    )
    arg_parser.add_argument(
        "--resume",
        type=str,
        default=None,
        help="Resume from a specific checkpoint file",
    )
    arg_parser.add_argument(
        "--compare-snapshot",
        type=str,
        default=None,
        help="Compare against a specific archived snapshot by date or date+time (e.g. 2026-04-08 or 2026-04-08_1805)",
    )
    arg_parser.add_argument(
        "--no-validation",
        action="store_true",
        help="Disable data validation and sanity checks",
    )
    arg_parser.add_argument(
        "--currency-report",
        action="store_true",
        help="Generate the standalone Currency Rate report via FZS commands",
    )
    arg_parser.add_argument(
        "--load-previous-rates",
        type=str,
        default=None,
        help="Path to a JSON/CSV/XLSX file to seed the previous-day rates (used with --previous-date)",
    )
    arg_parser.add_argument(
        "--previous-date",
        type=str,
        default=None,
        help="Effective date (YYYY-MM-DD) for --load-previous-rates; defaults to yesterday",
    )
    arg_parser.add_argument(
        "--user-token",
        type=str,
        default=None,
        help="Session token for the web dashboard account (overrides keyring and TRAVELPORT_USER_TOKEN env var)",
    )

    if prebuilt_args is not None:
        # Called from GUI with a ready-made Namespace - skip argparse entirely
        args = prebuilt_args
    elif len(sys.argv) == 1 and getattr(sys, "frozen", False):
        # Double-clicked exe with no flags - default to --auto
        args = arg_parser.parse_args(["--auto"])
    else:
        args = arg_parser.parse_args()

    use_tqdm = _should_use_tqdm(getattr(args, "_gui_mode", False))

    # CLI-only: install a global ESC watcher so users can abort without the GUI.
    # The GUI installs its own listener already, so skip when it drives the run.
    if not getattr(args, "_gui_mode", False):
        if _stop is None:
            import threading as _threading

            _stop = _threading.Event()
        _install_cli_esc_listener(_stop, logger)

    # QUICK PASTE MODE INTERCEPT
    if getattr(args, "quick_paste", False):
        # GUI mode: data was collected interactively before this function was
        # called and stored on args._quick_paste_data.  Skip all input() calls
        # so we never touch sys.stdin (which is None in the frozen exe).
        gui_qp = getattr(args, "_quick_paste_data", None)
        if gui_qp is not None:
            valid_commands = gui_qp["commands"]
            raw_texts = gui_qp["raw_texts"]
            raw_fs_texts = gui_qp["raw_fs_texts"]
        else:
            # CLI fallback - normal interactive flow
            from clipboard_util import clipboard_paste

            commands_input = input(
                "\n  Enter Commands (e.g., FDDACDOH/QR, FDDOHDAC/QR, FDDACMCT/WY): "
            ).strip()
            if not commands_input:
                print("  [!] Commands are required. Exiting.")
                sys.exit(1)

            raw_commands = [c.strip().upper() for c in commands_input.split(",")]

            valid_commands = []
            for c in raw_commands:
                parsed = parse_command(c)
                if parsed:
                    valid_commands.append(parsed)
                else:
                    print(f"  [!] Invalid command skipped: {c}")

            if not valid_commands:
                print("  [!] No valid commands provided. Exiting.")
                sys.exit(1)

            raw_texts = {}
            raw_fs_texts = {}

            for i, cmd in enumerate(valid_commands):
                airline = cmd["airline"]
                route = cmd["route"]
                print(f"\n" + "-" * 40)
                print(
                    f"  GATHERING DATA FOR {airline} {route} ({i+1}/{len(valid_commands)})"
                )
                print("-" * 40)

                input(
                    f"  [1/2] Please highlight and COPY the FD terminal output for {route}, then press ENTER..."
                )
                fd_text = clipboard_paste()
                if not fd_text or len(fd_text.strip()) < 10:
                    print(
                        "  [!] Clipboard seems empty or too short. Continuing anyway..."
                    )

                input(
                    f"  [2/2] Now, highlight and COPY the FS (Tax) terminal output for {route}, then press ENTER..."
                )
                fs_text = clipboard_paste()

                file_key = f"{airline}_{route}"
                raw_texts[file_key] = fd_text
                raw_fs_texts[file_key] = fs_text

        print("\n  Parsing manual data...")
        config = load_config(args.config)
        all_route_data = process_route_data(
            raw_texts,
            raw_fs_texts,
            config,
            enable_validation=False,
            show_progress=use_tqdm,
            stop_event=stop_event,
        )

        if not all_route_data:
            print("  [!] Failed to parse data. Ensure your copied text is valid.")
            if prebuilt_args is not None:
                return None  # GUI mode: don't sys.exit() in worker thread
            sys.exit(1)

        # Use first route for filename
        primary_route = f"{valid_commands[0]['airline']}_{valid_commands[0]['route'].replace('-', '')}"
        suffix = "etc" if len(valid_commands) > 1 else ""
        output_name = f"quick_report_{primary_route}_{suffix}_{datetime.now().strftime('%H%M')}.xlsx".replace(
            "__", "_"
        )
        output_path = os.path.join(REPORTS_DIR, output_name)

        result_path = generate_report(
            all_route_data, output_path, changes=None, config=config
        )

        # Record to Database
        record_to_database(all_route_data, config, mode="quick-paste")

        logger.info(f"\n  [OK] Successfully generated: {result_path}")
        try:
            os.startfile(result_path)
            logger.info("  Opening file automatically...")
        except Exception:
            pass
        if prebuilt_args is not None:
            return result_path  # GUI mode: return path instead of sys.exit()
        sys.exit(0)

    log_file = setup_logging()

    # Apply speed profile if specified (must be done before any automation imports)
    if args.speed:
        from constants import set_speed_profile

        set_speed_profile(args.speed)

    # Validate limit argument
    if args.limit:
        try:
            args.limit = validate_limit(args.limit)
        except ValidationError as e:
            logger.error(f"  {e}")
            sys.exit(1)

    logger.info("=" * 60)
    mode_label = "PENALTY" if args.penalty else ("TAX" if args.tax else "FARE")
    logger.info(f"  TRAVELPORT {mode_label} AUTOMATION TOOL")
    logger.info(f"  {datetime.now().strftime('%d-%b-%Y %H:%M')}")
    logger.info("=" * 60)

    # Show active speed profile
    from constants import ACTIVE_SPEED_PROFILE

    if args.speed:
        logger.info(f"  Speed Profile: {ACTIVE_SPEED_PROFILE.upper()} (CLI override)")
    else:
        logger.info(f"  Speed Profile: {ACTIVE_SPEED_PROFILE.upper()} (default)")
    logger.info("")

    try:
        # [1/4] Config
        logger.info("[1/4] Loading configuration...")
        config = load_config(args.config)
        logger.info("  [OK] Config loaded")
    except ConfigurationError as e:
        logger.error(f"  Configuration error: {e}")
        logger.error("  Please check your config.json file and try again.")
        sys.exit(1)
    except Exception as e:
        logger.error(f"  Unexpected error loading configuration: {e}")
        sys.exit(1)

    # Route Commands or Tax Airports
    commands = []
    tax_airports = {}
    if args.penalty and args.tax:
        logger.error("  Use either --tax or --penalty, not both together.")
        sys.exit(1)

    if args.tax:
        try:
            tax_airports, tax_selection = _select_tax_airports_for_run(config, args)
        except ValidationError as e:
            logger.error(f"  {e}")
            sys.exit(1)

        if not tax_airports and not tax_selection["requested_queries"]:
            logger.error("  No 'tax_airports' defined in config.")
            sys.exit(1)

        if tax_selection["resolutions"]:
            resolution_summary = "; ".join(
                (
                    f"{item['resolved_code']} ({item['display_name']}) "
                    f"via '{item['matched_alias']}' -> FTAX-{item['country_code']}"
                )
                for item in tax_selection["resolutions"]
            )
            logger.info(
                f"  [FILTER] Tax airports resolved: {resolution_summary}"
            )
        elif tax_selection["skipped_by_route_filter"] > 0:
            logger.info(
                f"  Filtered to {len(tax_airports)} airports matching configured routes "
                f"(skipped {tax_selection['skipped_by_route_filter']} unrelated)"
            )

        if args.limit > 0 and not tax_selection["resolutions"]:
            logger.info(f"  [TESTING] Limited to first {args.limit} airports")
        if tax_selection["resolutions"]:
            logger.info(f"  {len(tax_airports)} tax airport(s) selected for this run")
        else:
            logger.info(f"  {len(tax_airports)} tax airports loaded from config")
    else:
        import urllib.request as _ur
        from parser import load_commands_from_text

        commands_file = os.path.join(
            SCRIPT_DIR, config.get("commands_file", "commands.txt")
        )
        configured_commands: list[dict] = []

        if os.path.exists(commands_file):
            # User's local commands.txt - may have been customised; always prefer it.
            logger.info(f"  Loading local commands from {commands_file}")
            configured_commands = load_commands(commands_file)
        else:
            # First run - download the default command list from GitHub and save
            # it next to the exe so the user can edit it later.
            remote_url = config.get("commands_url", REMOTE_COMMANDS_URL)
            logger.info("  No commands.txt found - downloading defaults from remote...")
            try:
                req = _ur.Request(
                    remote_url, headers={"User-Agent": "TravelportAuto/1.0"}
                )
                with _ur.urlopen(req, timeout=10) as resp:
                    content = resp.read().decode("utf-8")
                with open(commands_file, "w", encoding="utf-8") as fh:
                    fh.write(content)
                configured_commands = load_commands_from_text(content)
                logger.info(
                    f"  Downloaded {len(configured_commands)} default commands -> saved to {commands_file}"
                )
                logger.info("  You can edit commands.txt to add or remove routes.")
            except Exception as exc:
                if args.route:
                    logger.warning(
                        f"  Could not download default commands: {exc}. Falling back to explicit route generation."
                    )
                else:
                    logger.error(f"  Could not download default commands: {exc}")
                    logger.error(
                        f"  Create a commands.txt file in {SCRIPT_DIR} with your FD commands."
                    )
                    sys.exit(1)

        commands = list(configured_commands)

        # Apply filters
        if commands:
            if args.route:
                routes = [
                    r.strip().upper().replace("-", "") for r in args.route.split(",")
                ]
                valid_routes = [r for r in routes if len(r) == 6]

                commands = [
                    c
                    for c in commands
                    if any(
                        _command_matches_route(c, rt, one_direction=args.one_direction)
                        for rt in valid_routes
                    )
                ]
                direction_scope = (
                    "exact direction only" if args.one_direction else "both directions"
                )
                logger.info(
                    f"  [FILTER] Limited to route(s) {args.route} ({direction_scope}): {len(commands)} commands remaining"
                )
            if args.airline:
                airlines = [a.strip().upper() for a in args.airline.split(",")]
                commands = [
                    c
                    for c in commands
                    if any(f"/{al}" in c["command"].upper() for al in airlines)
                ]
                logger.info(
                    f"  [FILTER] Limited to airline(s) {args.airline}: {len(commands)} commands remaining"
                )
            if args.limit > 0:
                commands = commands[: args.limit]
                logger.info(f"  [TESTING] Limited to first {args.limit} commands")
            if not args.route and args.limit == 0:
                logger.info(f"  {len(commands)} route commands loaded")

        explicit_routes: list[str] = []
        explicit_airlines: list[str] = []
        if args.route and (not commands or args.airline):
            # Two cases:
            #   1. No commands matched at all → generate the full explicit set (original fallback).
            #   2. Both --route and --airline are given but commands.txt only covers some of the
            #      (route, airline) pairs → generate the missing pairs so the user's airline
            #      list is fully honored. Previously these were silently dropped.
            try:
                generated, explicit_routes, explicit_airlines = _build_explicit_route_commands(
                    args.route,
                    args.airline,
                    config,
                    one_direction=args.one_direction,
                )
            except (ValidationError, ConfigurationError) as e:
                logger.error(f"  {e}")
                sys.exit(1)

            existing_cmds = {c["command"].upper() for c in commands}
            missing = [g for g in generated if g["command"].upper() not in existing_cmds]

            if not commands and generated:
                commands = generated
                direction_scope = (
                    "exact direction only" if args.one_direction else "both directions"
                )
                airline_scope = (
                    args.airline
                    if args.airline
                    else f"all configured airlines ({len(explicit_airlines)})"
                )
                logger.info(
                    "  [FALLBACK] No configured commands matched. "
                    f"Generated {len(commands)} command(s) for route(s) {', '.join(explicit_routes)} "
                    f"using {airline_scope} ({direction_scope})"
                )
            elif missing and args.airline:
                commands = commands + missing
                logger.info(
                    f"  [FILL] Generated {len(missing)} command(s) for (route, airline) pairs "
                    f"missing from commands.txt: {', '.join(m['command'] for m in missing)}"
                )

            if commands and args.limit > 0:
                commands = commands[: args.limit]
                logger.info(f"  [TESTING] Limited to first {args.limit} commands")

        if not commands:
            if args.route:
                logger.error(
                    f"  No commands available for route(s) {args.route}. "
                    "Provide --airline or add matching commands to commands.txt."
                )
            else:
                logger.error(
                    "  No valid route commands are configured. Please update commands.txt and try again."
                )
            sys.exit(1)
    logger.info("")

    # Initialize checkpoint manager if enabled
    checkpoint_mgr = None
    enable_validation = not args.no_validation

    if args.checkpoint or args.resume:
        os.makedirs(CHECKPOINT_DIR, exist_ok=True)

        if args.resume:
            # Resume from specific checkpoint
            checkpoint_mgr = CheckpointManager(
                CHECKPOINT_DIR,
                session_name=os.path.basename(args.resume)
                .replace("checkpoint_", "")
                .replace(".json", ""),
            )
            if checkpoint_mgr.load_checkpoint():
                logger.info(f"  [CHECKPOINT] Resuming from: {args.resume}")
                stats = checkpoint_mgr.get_progress_stats()
                logger.info(
                    f"  [CHECKPOINT] Previously completed: {stats['completed']} commands"
                )
            else:
                logger.warning(
                    f"  [CHECKPOINT] Could not load checkpoint: {args.resume}"
                )
                logger.info("  [CHECKPOINT] Starting fresh")
        else:
            # New checkpoint session
            checkpoint_mgr = CheckpointManager(CHECKPOINT_DIR)
            logger.info(f"  [CHECKPOINT] Checkpoint mode enabled")
            logger.info(f"  [CHECKPOINT] Session: {checkpoint_mgr.session_name}")

    def _stop_run(message: str, partial_data: dict | None = None):
        logger.info(message)
        if checkpoint_mgr:
            checkpoint_mgr.save_checkpoint()
            logger.info(
                f"  [CHECKPOINT] Final checkpoint saved: {len(checkpoint_mgr.completed_commands)} completed"
            )

        # Generate partial report with whatever data was collected
        if partial_data:
            try:
                timestamp_full = datetime.now().strftime("%Y-%m-%d_%H%M")
                if getattr(args, "tax", False):
                    partial_path = os.path.join(
                        REPORTS_DIR, f"tax_report_{timestamp_full}_partial.xlsx"
                    )
                    from tax_report import generate_tax_report

                    result = generate_tax_report(partial_data, partial_path, None, config)
                else:
                    partial_path = os.path.join(
                        REPORTS_DIR, f"fare_report_{timestamp_full}_partial.xlsx"
                    )
                    result = generate_report(
                        partial_data, partial_path, None, config,
                        only_currency=getattr(args, "only_currency", False),
                    )
                logger.info(f"  [PARTIAL] Partial report saved: {result}")
                return result
            except Exception as exc:
                logger.warning(f"  [PARTIAL] Could not generate partial report: {exc}")

        return None

    # [2/4] Extraction
    failed_commands = []

    # CURRENCY RATE REPORT MODE
    if getattr(args, "currency_report", False):
        result_path = _run_currency_report_mode(args, config, _stop)
        elapsed = _time.time() - start_time
        minutes = int(elapsed // 60)
        seconds = int(elapsed % 60)
        logger.info("")
        logger.info("=" * 60)
        logger.info("  CURRENCY REPORT SUMMARY")
        logger.info("-" * 60)
        logger.info(f"  Report:   {result_path}")
        logger.info(f"  Duration: {minutes}m {seconds}s")
        logger.info("=" * 60)
        if prebuilt_args is not None:
            return result_path
        return result_path

    # PENALTY MODE EXTRACTION
    if args.penalty:
        penalty_records = []

        if not args.auto:
            logger.error("  Penalty mode currently requires --auto.")
            sys.exit(1)

        logger.info("[2/3] AUTO MODE: Connecting to Smartpoint UI for penalties...")
        if not commands:
            logger.error("  No commands found to auto-run.")
            sys.exit(1)

        from smartpoint_automation import SmartpointAutomation, StopRequested

        automation = SmartpointAutomation(stop_event=_stop)
        try:
            if not automation.connect():
                logger.error("  Please ensure Smartpoint is open and the title matches.")
                sys.exit(1)

            username, password, pcc = CredentialManager.get_credentials()
            if username and password:
                logger.info("  Credentials loaded from environment")
                if _env_truthy("SMARTPOINT_SKIP_AUTO_LOGIN"):
                    logger.info(
                        "  Automatic login skipped because SMARTPOINT_SKIP_AUTO_LOGIN is enabled."
                    )
                else:
                    try:
                        if not automation.login(username, password, pcc):
                            logger.error(
                                "  Automatic login did not complete. Please sign in manually and try again."
                            )
                            sys.exit(1)
                    except Exception as e:
                        logger.error(f"  Login failed: {e}")
                        sys.exit(1)
            else:
                logger.info("  No credentials found - continuing without login")

            automation.refresh_terminal()
            os.makedirs(RAW_PENALTY_DIR, exist_ok=True)

            if use_tqdm:
                command_iter = tqdm(
                    commands, desc="Extracting penalties", unit="cmd", file=_tqdm_stream()
                )
            else:
                command_iter = commands
                logger.info(f"  Executing {len(commands)} commands...")

            penalty_loop_start = _time.time()
            penalty_total = len(commands)
            for i, cmd in enumerate(command_iter, 1):
                if _stop and _stop.is_set():
                    logger.info("  [STOP] Stop requested - finishing after this point.")
                    break
                cmd_str = cmd["command"]
                if not use_tqdm:
                    logger.info(f"  [{i}/{penalty_total}] {cmd_str}")
                    _log_eta(
                        logger,
                        penalty_loop_start,
                        i,
                        penalty_total,
                        label="penalty commands",
                        every=5,
                    )

                command_penalties = automation.run_penalty_command(cmd_str)
                if not command_penalties:
                    failed_commands.append(cmd_str)
                    logger.warning(f"    [!] No penalty popups captured for {cmd_str}")
                    continue

                file_key = generate_file_key(cmd)
                for penalty_capture in command_penalties:
                    fare_basis = penalty_capture.get("fare_basis")
                    journey_type = "RT" if penalty_capture.get("is_rt") else "OW"
                    parsed_penalty = parse_penalty_text(
                        penalty_capture.get("raw_penalty_text", ""),
                        airline=penalty_capture.get("airline", cmd.get("airline")),
                        route=cmd.get("route"),
                        fare_basis=fare_basis,
                        rbd=penalty_capture.get("rbd"),
                        fare_amount=penalty_capture.get("fare"),
                        journey_type=journey_type,
                    )

                    if not parsed_penalty.get("rules"):
                        logger.warning(
                            f"    [!] No structured penalty rules parsed for {fare_basis}"
                        )

                    penalty_records.append(parsed_penalty)

                    backup_name = f"{file_key}_{journey_type}_{_safe_filename(fare_basis)}_penalty.txt"
                    backup_path = os.path.join(RAW_PENALTY_DIR, backup_name)
                    try:
                        with open(backup_path, "w", encoding="utf-8") as f:
                            f.write(parsed_penalty.get("raw_penalty_text", ""))
                    except Exception as e:
                        logger.warning(f"    Could not save penalty backup: {e}")

            automation.show_completion_signal()
            logger.info("")
        except StopRequested:
            return _stop_run("  [STOP] Stop requested - penalty extraction stopped.")

        if _stop and _stop.is_set():
            return _stop_run("  [STOP] Stop requested - skipping penalty report generation.")

        logger.info("[3/3] Generating penalty report...")
        timestamp_full = datetime.now().strftime("%Y-%m-%d_%H%M")
        output_path = (
            args.output
            if args.output
            else os.path.join(REPORTS_DIR, f"penalty_report_{timestamp_full}.xlsx")
        )
        result_path = generate_penalty_report(penalty_records, output_path)
        try:
            from usage_tracker import increment as _inc
            _inc("penalty_runs")
        except Exception:
            pass

        elapsed = _time.time() - start_time
        minutes = int(elapsed // 60)
        seconds = int(elapsed % 60)

        logger.info("")
        logger.info("=" * 60)
        logger.info("  RUN SUMMARY")
        logger.info("-" * 60)
        logger.info(f"  Report:   {result_path}")
        logger.info(f"  Items:    {len(penalty_records)} fare basis records parsed")
        if failed_commands:
            logger.info(
                f"  Failed:   {len(failed_commands)} ({', '.join(failed_commands)})"
            )
        else:
            logger.info("  Failed:   None")
        logger.info(f"  Changes:  Skipped")
        logger.info(f"  Duration: {minutes}m {seconds}s")
        logger.info(f"  Log:      {log_file}")
        logger.info("=" * 60)

        return result_path

    # TAX MODE EXTRACTION
    if args.tax:
        tax_data = {}  # airport_code -> {'taxes': [ detail_data... ]}

        if args.auto:
            logger.info("[2/4] AUTO MODE: Connecting to Smartpoint UI for FTAX...")
            from smartpoint_automation import SmartpointAutomation, StopRequested

            automation = SmartpointAutomation(stop_event=_stop)
            try:
                if not automation.connect():
                    logger.error("  Please ensure Smartpoint is open.")
                    sys.exit(1)

                automation.refresh_terminal()
                from tax_parser import parse_ftax_list, parse_ftax_detail

                # Use tqdm for progress if available
                airport_items = tax_airports.items()
                if use_tqdm:
                    airport_items = tqdm(
                        list(airport_items),
                        desc="Extracting tax data",
                        unit="airport",
                        file=_tqdm_stream(),
                    )
                else:
                    airport_items = list(airport_items)

                tax_loop_start = _time.time()
                tax_total = len(tax_airports)
                for index, (airport_code, airport_info) in enumerate(airport_items, 1):
                    if _stop and _stop.is_set():
                        logger.info("  [STOP] Stop requested - finishing after this point.")
                        break
                    country_code = airport_info["country"]
                    display_name = _tax_airport_display_name(
                        airport_code, airport_info, config
                    )

                    # Only log if not using tqdm
                    if not use_tqdm:
                        logger.info(
                            f"  [{index}/{tax_total}] Airport: {airport_code} ({display_name}) -> FTAX-{country_code}"
                        )
                        _log_eta(
                            logger,
                            tax_loop_start,
                            index,
                            tax_total,
                            label="tax airports",
                            every=5,
                        )

                    # Get tax types
                    list_cmd = f"FTAX-{country_code}"
                    list_text = automation.run_command(list_cmd, max_pages=1)

                    # Debug backup
                    os.makedirs(RAW_DATA_DIR, exist_ok=True)
                    backup_path = os.path.join(RAW_DATA_DIR, f"{list_cmd}.txt")
                    with open(backup_path, "w", encoding="utf-8") as f:
                        f.write(list_text)

                    tax_types = parse_ftax_list(list_text)
                    logger.info(
                        f"    Found {len(tax_types)} tax types: {', '.join(t['code'] for t in tax_types)}"
                    )

                    if not tax_types:
                        logger.warning(
                            f"    [DEBUG] Raw list text ({len(list_text)} chars): {list_text[:300]}"
                        )

                    airport_tax_details = []
                    for idx, t in enumerate(tax_types, 1):
                        detail_text = automation.run_ftax_command(
                            country_code, t["code"], tax_index=idx
                        )

                        # Save raw detail text for debugging
                        detail_backup = os.path.join(
                            RAW_DATA_DIR, f"FTAX-{country_code}_{t['code']}.txt"
                        )
                        with open(detail_backup, "w", encoding="utf-8") as f:
                            f.write(detail_text)

                        if not detail_text or len(detail_text.strip()) < 20:
                            failed_commands.append(f"{country_code}/{t['code']}")
                            logger.warning(f"    Failed to extract details for {t['code']}")
                            continue

                        logger.debug(
                            f"      [DEBUG] Raw detail text first 200 chars: {detail_text[:200]}"
                        )

                        detail_data = parse_ftax_detail(detail_text, t["code"], t["name"])
                        airport_tax_details.append(detail_data)
                        logger.info(
                            f"      {t['code']} -> {len(detail_data['sections'])} sections extracted."
                        )

                        if not detail_data["sections"]:
                            logger.warning(
                                f"      [DEBUG] 0 sections! Full text ({len(detail_text)} chars): {detail_text[:500]}"
                            )

                        # Return to tax list for next tax type
                        if idx < len(tax_types):
                            automation.return_to_tax_list(country_code)

                    tax_data[airport_code] = {
                        "taxes": airport_tax_details,
                        "_country": country_code,
                    }

                automation.show_completion_signal()
            except StopRequested:
                return _stop_run("  [STOP] Stop requested - tax extraction stopped.", partial_data=tax_data if tax_data else None)
        else:
            logger.error("  Manual loading of taxes not implemented. Use --auto.")
            sys.exit(1)

        all_route_data = tax_data  # Alias for reporting
        if _stop and _stop.is_set():
            return _stop_run("  [STOP] Stop requested - skipping tax report generation.", partial_data=all_route_data)
        logger.info("")

    # FARE MODE EXTRACTION
    else:
        raw_texts = {}
        raw_fs_texts = {}
        if args.auto:
            logger.info("[2/4] AUTO MODE: Connecting to Smartpoint UI...")
            if not commands:
                logger.error("  No commands found to auto-run.")
                sys.exit(1)

            from smartpoint_automation import SmartpointAutomation

            automation = SmartpointAutomation(stop_event=_stop)

            if not automation.connect():
                logger.error(
                    "  Please ensure Smartpoint is open and the title matches."
                )
                sys.exit(1)

            # Try to get credentials from environment
            username, password, pcc = CredentialManager.get_credentials()
            if username and password:
                logger.info("  Credentials loaded from environment")
                if _env_truthy("SMARTPOINT_SKIP_AUTO_LOGIN"):
                    logger.info(
                        "  Automatic login skipped because SMARTPOINT_SKIP_AUTO_LOGIN is enabled."
                    )
                else:
                    try:
                        if not automation.login(username, password, pcc):
                            logger.error(
                                "  Automatic login did not complete. Please sign in manually and try again."
                            )
                            sys.exit(1)
                    except Exception as e:
                        logger.error(f"  Login failed: {e}")
                        sys.exit(1)
            else:
                logger.info("  No credentials found - continuing without login")
                logger.info("  To enable automatic login, set environment variables:")
                logger.info(
                    "    SMARTPOINT_USERNAME, SMARTPOINT_PASSWORD, SMARTPOINT_PCC"
                )

            logger.debug("  Initializing terminal state...")
            automation.refresh_terminal()

            # Filter commands if resuming from checkpoint
            if checkpoint_mgr:
                original_count = len(commands)
                commands = checkpoint_mgr.get_remaining_commands(commands)
                skipped = original_count - len(commands)
                if skipped > 0:
                    logger.info(
                        f"  [CHECKPOINT] Skipping {skipped} already completed commands"
                    )
                    logger.info(f"  [CHECKPOINT] {len(commands)} commands remaining")

            MAX_RETRIES = MAX_RETRIES_COMMAND

            # Use tqdm for progress if available
            if use_tqdm and not checkpoint_mgr:
                # Use simple progress bar
                command_iter = tqdm(
                    commands,
                    desc="Executing commands",
                    unit="cmd",
                    file=_tqdm_stream(),
                )
            elif use_tqdm and checkpoint_mgr:
                # Use progress bar with initial progress
                command_iter = tqdm(
                    commands,
                    desc="Executing commands",
                    unit="cmd",
                    initial=len(checkpoint_mgr.completed_commands),
                    total=len(checkpoint_mgr.completed_commands) + len(commands),
                    file=_tqdm_stream(),
                )
            else:
                # No progress bar
                command_iter = commands
                logger.info(f"  Executing {len(commands)} commands...")

            commands_attempted = 0
            fare_loop_start = _time.time()
            fare_total = len(commands)
            for i, cmd in enumerate(command_iter, 1):
                if _stop and _stop.is_set():
                    logger.info("  [STOP] Stop requested - finishing after this point.")
                    break
                cmd_str = cmd["command"]

                # Skip if already completed (double-check in case of concurrent runs)
                if checkpoint_mgr and checkpoint_mgr.is_completed(cmd_str):
                    continue

                commands_attempted += 1

                # Only log if not using tqdm
                if not use_tqdm:
                    logger.info(f"  [{i}/{fare_total}] {cmd_str}")
                    _log_eta(
                        logger,
                        fare_loop_start,
                        i,
                        fare_total,
                        label="fare commands",
                        every=5,
                    )

                terminal_text = ""
                file_key = generate_file_key(cmd)
                fd_no_fares = False

                # -- FD EXTRACTION --
                if not args.only_yq and not args.only_currency:
                    for attempt in range(1, MAX_RETRIES + 1):
                        try:
                            terminal_text = automation.run_command(cmd["command"])
                            if _fd_output_is_no_fares(terminal_text):
                                fd_no_fares = True
                                logger.info(
                                    "    [SKIP] Smartpoint returned no fares; not retrying this FD command."
                                )
                                break
                            if terminal_text and len(terminal_text.strip()) > 50:
                                break
                            logger.warning(
                                f"    Attempt {attempt}: insufficient data ({len(terminal_text)} chars)"
                            )
                        except Exception as e:
                            logger.warning(f"    Attempt {attempt} failed: {e}")

                        if attempt < MAX_RETRIES:
                            delay = constants.RETRY_DELAY
                            logger.info(f"    Retrying in {delay}s...")
                            _time.sleep(delay)
                            automation.refresh_terminal()

                    if terminal_text and len(terminal_text.strip()) > 50:
                        raw_texts[file_key] = terminal_text
                        logger.info(
                            f"    [OK] Fare data captured ({len(terminal_text)} chars)"
                        )

                        backup_path = os.path.join(RAW_DATA_DIR, f"{file_key}.txt")
                        os.makedirs(os.path.dirname(backup_path) or ".", exist_ok=True)
                        try:
                            with open(backup_path, "w", encoding="utf-8") as f:
                                f.write(terminal_text)
                            logger.debug(f"    Backup saved: {backup_path}")
                        except Exception as e:
                            logger.warning(f"    Could not save backup: {e}")
                    elif fd_no_fares:
                        logger.info(
                            f"    [SKIP] No fare rows available for {cmd['command']}."
                        )
                    else:
                        failed_commands.append(cmd["command"])
                        logger.error(
                            f"    [FAILED] FAILED after {MAX_RETRIES} attempts: {cmd['command']}"
                        )
                        logger.error(
                            f"    Final data length: {len(terminal_text) if terminal_text else 0} chars"
                        )
                        logger.error(f"    This command will be skipped in the report")

                        # Mark as failed in checkpoint
                        if checkpoint_mgr:
                            checkpoint_mgr.mark_failed(cmd_str)
                else:
                    logger.info(
                        "    [SKIP] Skipping FD extraction (--only-yq/--only-currency)"
                    )

                # -- FS EXTRACTION --
                if _should_run_fs_extraction(args, terminal_text):
                    base_cmd = cmd["command"].strip()
                    if (
                        len(base_cmd) >= 11
                        and base_cmd.startswith("FD")
                        and "/" in base_cmd
                        and len(base_cmd.split("/")[0]) == 8
                    ):
                        src = base_cmd[2:5]
                        dst = base_cmd[5:8]
                        airline = base_cmd.split("/")[1][:2]

                        fs_expanded = ""
                        # Two-window fallback: 7 consecutive days ~1 month out,
                        # then 7 more ~3 months out if the first window yields
                        # no pure-airline option.
                        fs_date_offsets: list[int] = []
                        for window_start in (
                            FS_DATE_OFFSET_START,
                            FS_DATE_FALLBACK_OFFSET,
                        ):
                            fs_date_offsets.extend(
                                range(
                                    window_start,
                                    window_start + FS_DATE_WINDOW_DAYS,
                                    FS_DATE_STEP,
                                )
                            )

                        for fs_date_offset in fs_date_offsets:
                            date_str = (
                                (datetime.now() + timedelta(days=fs_date_offset))
                                .strftime("%d%b")
                                .upper()
                            )
                            logger.info(
                                f"    [FS Checkout] Extracting tax details for {src}-{dst} on {date_str}..."
                            )

                            fs_result = automation.run_fs_command(
                                src, dst, date_str, airline
                            )
                            fs_no_results = _fs_output_is_no_results(fs_result)

                            # Freshness check: poll until the terminal shows the
                            # expected date, rather than sleeping a fixed 1.5s.
                            if (
                                not fs_no_results
                                and date_str.upper() not in fs_result.upper()
                            ):
                                logger.warning(
                                    f"      [!] Screen hasn't updated to {date_str} yet, polling..."
                                )
                                fs_result = automation._wait_for_response(
                                    fs_result,
                                    timeout=constants.RETRY_DELAY * 2,
                                    min_wait=0.2,
                                    poll_interval=0.15,
                                    stability_checks=1,
                                )
                                fs_no_results = _fs_output_is_no_results(fs_result)
                            elif fs_no_results:
                                logger.warning(
                                    f"      [!] FS returned no usable result for {date_str}; not polling this screen."
                                )

                            # Log raw results for diagnostics
                            with open(os.path.join(LOG_DIR, "fs_debug.log"), "a", encoding="utf-8") as f:
                                f.write(
                                    f"\n--- {date_str} {src}-{dst} /{airline} ---\n"
                                )
                                f.write(fs_result)
                                f.write("\n" + "=" * 50 + "\n")

                            current_fs_page = fs_result
                            target_option_index = None
                            target_option_number = None
                            fs_page_number = 1
                            max_fs_pages = 5
                            rechecked_current_fs_page = False

                            while True:
                                (
                                    target_option_index,
                                    target_option_number,
                                    option_count,
                                ) = _find_pure_airline_option_in_fs_page(
                                    current_fs_page, airline
                                )

                                logger.info(
                                    f"      [DEBUG] Parsing {option_count} options for {airline} on FS page {fs_page_number}..."
                                )

                                if target_option_index is not None:
                                    logger.info(
                                        f"      [OK] Pure {airline} itinerary found in Option {target_option_number} on FS page {fs_page_number}."
                                    )
                                    fs_result = current_fs_page
                                    break

                                if (
                                    not rechecked_current_fs_page
                                    and not _fs_output_is_no_results(current_fs_page)
                                ):
                                    settled_fs_page = (
                                        automation._wait_for_stable_screen(
                                            max_polls=4, interval=0.25
                                        )
                                    )
                                    rechecked_current_fs_page = True

                                    if _should_recheck_same_fs_page(
                                        current_fs_page, settled_fs_page, airline
                                    ):
                                        logger.info(
                                            f"      [DEBUG] Rechecking FS page {fs_page_number} for {airline} after additional settle..."
                                        )
                                        current_fs_page = settled_fs_page
                                        continue

                                if option_count == 0:
                                    if _fs_output_is_no_results(current_fs_page):
                                        logger.warning(
                                            f"      [!] No valid FS results for {date_str}. Trying next date..."
                                        )
                                        break

                                    logger.warning(
                                        f"      [!] Waiting for terminal content (offset {fs_date_offset})..."
                                    )
                                    break

                                if (
                                    fs_page_number >= max_fs_pages
                                    or automation._has_end_signal(current_fs_page)
                                ):
                                    logger.warning(
                                        f"      [!] No pure {airline} options found on {date_str} after {fs_page_number} FS page(s)."
                                    )
                                    break

                                logger.info(
                                    f"      [DEBUG] No pure {airline} option on FS page {fs_page_number}; checking next FS page on the same date..."
                                )

                                if automation.click_more_prompt_link(current_fs_page):
                                    next_fs_page = automation._wait_for_response(
                                        current_fs_page,
                                        timeout=constants.COMMAND_WAIT_MEDIUM + 0.5,
                                        min_wait=0.0,
                                        stability_checks=1,
                                    )
                                    next_fs_page = automation._wait_for_stable_screen(
                                        max_polls=3, interval=0.2
                                    )
                                else:
                                    logger.warning(
                                        f"      [!] More Flights was not clickable on {date_str}; staying off MD and trying next date."
                                    )
                                    break

                                if (
                                    not next_fs_page
                                    or automation._has_invalid(next_fs_page)
                                    or next_fs_page.strip() == current_fs_page.strip()
                                ):
                                    logger.warning(
                                        f"      [!] Could not advance FS pagination on {date_str}; trying next date."
                                    )
                                    break

                                current_fs_page = next_fs_page
                                fs_page_number += 1
                                rechecked_current_fs_page = False

                            if target_option_index is None:
                                _time.sleep(0.5)
                                continue

                            # Click the D button using text-to-coordinate mapping
                            fs_expanded = automation.click_d_button(
                                target_option_index, fs_result
                            )

                            # Validate: accept both classic fare/tax lines and airline-
                            # specific detail screens that still parse into usable tax data.
                            if fs_expanded and looks_like_fs_tax_breakdown(fs_expanded):
                                logger.info(
                                    f"      [OK] Tax breakdown extracted via D-click"
                                )
                            else:
                                logger.warning(
                                    f"      [!] D-click did not return expected tax data."
                                )
                                fs_expanded = ""
                            break  # Exit the date-stepping for loop

                        if fs_expanded and len(fs_expanded.strip()) > 50:
                            raw_fs_texts[file_key] = fs_expanded
                            fs_backup_path = os.path.join(
                                RAW_DATA_DIR, f"{file_key}_FS.txt"
                            )
                            try:
                                with open(fs_backup_path, "w", encoding="utf-8") as f:
                                    f.write(fs_expanded)
                            except Exception:
                                pass
                else:
                    logger.info(
                        "    [SKIP] Skipping FS extraction because FD returned no fare data."
                    )

                # Mark command as completed and save checkpoint
                if checkpoint_mgr and cmd_str not in failed_commands:
                    checkpoint_mgr.mark_completed(cmd_str)
                    # Save checkpoint every 10 commands to avoid excessive I/O
                    if len(checkpoint_mgr.completed_commands) % 10 == 0:
                        checkpoint_mgr.save_checkpoint()

            # Final checkpoint save
            if checkpoint_mgr:
                checkpoint_mgr.save_checkpoint()
                logger.info(
                    f"  [CHECKPOINT] Final checkpoint saved: {len(checkpoint_mgr.completed_commands)} completed"
                )

            # -- MERGED FTAX EXTRACTION --
            ftax_data = None
            if getattr(args, "include_ftax", False):
                logger.info(
                    "\n  --include-ftax: Extracting FTAX data for involved airports..."
                )
                unique_airports = set()
                for c in commands:
                    if "-" in c.get("route", ""):
                        orig, dest = c["route"].split("-")
                        unique_airports.add(orig)
                        unique_airports.add(dest)

                tax_airports_config = config.get("tax_airports", {})
                target_airports = {}
                for a in unique_airports:
                    if a in tax_airports_config:
                        target_airports[a] = tax_airports_config[a]
                    else:
                        logger.warning(
                            f"    Skipping FTAX for {a}: missing in config.json 'tax_airports'"
                        )

                if target_airports:
                    from tax_parser import parse_ftax_list, parse_ftax_detail

                    ftax_data = {}

                    logger.info(
                        f"  Found {len(target_airports)} airports for FTAX profiling: {', '.join(target_airports.keys())}"
                    )

                    for index, (acode, ainfo) in enumerate(target_airports.items(), 1):
                        ccode = ainfo["country"]
                        display_name = _tax_airport_display_name(acode, ainfo, config)
                        logger.info(
                            f"  [{index}/{len(target_airports)}] Airport: {acode} ({display_name}) -> FTAX-{ccode}"
                        )

                        list_cmd = f"FTAX-{ccode}"
                        list_text = automation.run_command(list_cmd, max_pages=1)
                        tax_types = parse_ftax_list(list_text)
                        logger.info(
                            f"    Found {len(tax_types)} tax types: {', '.join(t['code'] for t in tax_types)}"
                        )

                        airport_tax_details = []
                        for idx, t in enumerate(tax_types, 1):
                            detail_text = automation.run_ftax_command(
                                ccode, t["code"], tax_index=idx
                            )
                            if not detail_text or len(detail_text.strip()) < 20:
                                logger.warning(
                                    f"    Failed to extract details for {t['code']}"
                                )
                                continue

                            detail_data = parse_ftax_detail(
                                detail_text, t["code"], t["name"]
                            )
                            airport_tax_details.append(detail_data)
                            logger.info(
                                f"      {t['code']} -> {len(detail_data['sections'])} sections extracted."
                            )

                            if idx < len(tax_types):
                                automation.return_to_tax_list(ccode)

                        ftax_data[acode] = {"taxes": airport_tax_details}

            automation.show_completion_signal()

            # Show execution summary
            total_commands = len(commands)
            successful_commands = commands_attempted - len(failed_commands)
            skipped_commands = total_commands - commands_attempted
            logger.info("")
            logger.info("=" * 60)
            logger.info("  EXECUTION SUMMARY")
            logger.info("=" * 60)
            logger.info(f"  Total commands: {total_commands}")
            logger.info(f"  Attempted: {commands_attempted}")
            logger.info(f"  Successful: {successful_commands}")
            logger.info(f"  Failed: {len(failed_commands)}")
            if skipped_commands > 0:
                logger.info(f"  Skipped (stopped early): {skipped_commands}")
            if failed_commands:
                logger.warning("  Failed commands:")
                for fc in failed_commands:
                    logger.warning(f"    - {fc}")
                logger.warning("  Note: Failed commands will not appear in the report")
            elif commands_attempted == total_commands:
                logger.info("  [OK] All commands completed successfully!")
            logger.info("=" * 60)

            if _stop and _stop.is_set():
                if raw_texts:
                    logger.info(
                        "  [STOP] Stop requested - generating partial report from captured fare data."
                    )
                else:
                    return _stop_run(
                        "  [STOP] Stop requested - no captured fare data to report."
                    )

        else:
            logger.info("[2/4] MANUAL MODE: Loading raw GDS data from disk...")
            raw_texts, raw_fs_texts = load_raw_data(RAW_DATA_DIR)
            if not raw_texts:
                show_usage()
                sys.exit(1)
            logger.info(
                f"  Loaded {len(raw_texts)} fare file(s) and {len(raw_fs_texts)} tax detail file(s)"
            )
        logger.info("")

        # Parse Fares
        logger.info("[3/4] Parsing fare and tax data...")
        all_route_data = process_route_data(
            raw_texts,
            raw_fs_texts,
            config,
            enable_validation,
            show_progress=use_tqdm,
            stop_event=_stop,
        )
        if _stop and _stop.is_set():
            return _stop_run("  [STOP] Stop requested - parsing stopped.", partial_data=all_route_data)
        if not all_route_data:
            logger.error("  No fare data could be parsed.")
            sys.exit(1)
        logger.info("")

    # [3.5] Change detection (for both modes)
    changes = None
    if _stop and _stop.is_set():
        return _stop_run("  [STOP] Stop requested - skipping change detection.", partial_data=all_route_data)
    if not args.no_changes:
        logger.info("[3.5] Detecting changes...")
        archive_subdir = "tax" if args.tax else "fare"
        archive_path = os.path.join(ARCHIVE_DIR, archive_subdir)
        os.makedirs(archive_path, exist_ok=True)

        latest_snapshot_data = load_latest_snapshot(archive_path)
        previous_data = latest_snapshot_data
        comparison_snapshot_id = None
        using_db_snapshot = False

        # Try to load previous snapshot from DB first
        if DatabaseManager and _resolve_database_url(config):
            try:
                _db_check = DatabaseManager(_resolve_database_url(config))
                if _db_check.connect():
                    _run_mode_for_check = "tax-mode" if args.tax else None
                    prev_run_id = _db_check.get_previous_run_id(
                        run_mode=_run_mode_for_check
                    )
                    if prev_run_id > 0:
                        if args.tax:
                            db_snapshot = _db_check.load_tax_snapshot(prev_run_id)
                        else:
                            db_snapshot = _db_check.load_fare_snapshot(prev_run_id)
                        if db_snapshot:
                            previous_data = db_snapshot
                            using_db_snapshot = True
                            logger.info(f"  Using DB snapshot from run {prev_run_id}")
                    _db_check.close()
            except Exception as _db_snap_err:
                logger.debug("  DB snapshot load failed: %s", _db_snap_err)

        if not previous_data:
            # Fall back to JSON archive
            previous_data = latest_snapshot_data
            if previous_data:
                logger.info("  Using JSON archive snapshot (no DB history yet)")

        if args.compare_snapshot:
            try:
                previous_data, comparison_snapshot_id = load_snapshot_by_reference(
                    archive_path, args.compare_snapshot
                )
            except ValueError as e:
                logger.error(f"  {e}")
                sys.exit(1)

            if previous_data is None:
                logger.error(
                    f"  No archived snapshot found for '{args.compare_snapshot}'."
                )
                sys.exit(1)

            logger.info(f"  Comparing against snapshot: {comparison_snapshot_id}")

        if previous_data:
            if args.tax:
                changes = detect_tax_changes(all_route_data, previous_data)
                if changes and any(changes.values()):
                    logger.info(format_tax_change_summary(changes))
                else:
                    logger.info("  No changes from previous tax data.")
            else:
                changes = detect_changes(all_route_data, previous_data)
                if changes and any(changes.values()):
                    logger.info(format_change_summary(changes))
                else:
                    logger.info("  No changes from previous data.")
        else:
            logger.info("  No previous data found. First run - baseline saved.")

        # Only write JSON snapshot when NOT using DB as primary source
        if not using_db_snapshot:
            if snapshot_has_changed(all_route_data, latest_snapshot_data):
                ts = datetime.now().strftime("%Y-%m-%d_%H%M")
                save_snapshot(all_route_data, archive_path, date_str=ts)
            else:
                logger.info("  Snapshot unchanged; skipping archive write.")
        else:
            logger.info("  DB is primary snapshot source; skipping JSON archive write.")
        logger.info("")

    # [4/4] Generate Report
    if _stop and _stop.is_set():
        return _stop_run("  [STOP] Stop requested - skipping report generation.", partial_data=all_route_data)
    logger.info("[4/4] Generating Excel report...")

    if args.output:
        output_path = args.output
    else:
        if args.tax:
            timestamp_full = datetime.now().strftime("%Y-%m-%d_%H%M")
            output_name = f"tax_report_{timestamp_full}.xlsx"
            output_path = os.path.join(REPORTS_DIR, output_name)
        else:
            timestamp_full = datetime.now().strftime("%Y-%m-%d_%H%M")
            output_filename = config.get(
                "excel_output_filename", "fare_report_{date}.xlsx"
            )
            output_filename = output_filename.replace("{date}", timestamp_full)
            output_path = os.path.join(REPORTS_DIR, output_filename)

    if args.tax:
        from tax_report import generate_tax_report

        result_path = generate_tax_report(all_route_data, output_path, changes, config)
    else:
        # Pass only_currency flag down to specifically skip the main sheets if needed
        result_path = generate_report(
            all_route_data,
            output_path,
            changes,
            config,
            only_currency=args.only_currency,
        )

        # Merge FTAX if enabled
        if (
            getattr(args, "include_ftax", False)
            and "ftax_data" in locals()
            and ftax_data
        ):
            logger.info("  Appending FTAX sheets to Fare Report...")
            import openpyxl
            from tax_report import _build_summary_sheet, _build_details_sheet

            try:
                wb = openpyxl.load_workbook(result_path)
                ws_summary = wb.create_sheet("FTAX Summary")
                ws_details = wb.create_sheet("FTAX Detailed Rates")
                _build_summary_sheet(ws_summary, ftax_data, config)
                _build_details_sheet(ws_details, ftax_data, config)
                wb.save(result_path)
                logger.info(f"  [OK] Attached FTAX to {result_path}")
            except Exception as e:
                logger.error(f"  Failed to append FTAX sheets: {e}")

    # [DB] Optional persistence - keep current file/report flow unchanged
    _run_mode = "auto" if not args.tax else "tax-mode"
    _db_run_id = record_to_database(all_route_data, config, mode=_run_mode)
    try:
        from usage_tracker import increment as _inc
        _route_count = len(all_route_data) if all_route_data else 1
        if args.tax:
            _inc("ftax_airports", _route_count)
        else:
            _inc("fare_routes", _route_count)
    except Exception:
        pass

    # [BQ] Push to BigQuery if configured
    if _bq_pusher and _bq_pusher.is_configured():
        _run_time = datetime.now()
        try:
            if args.tax:
                _bq_pusher.push_tax_snapshot(all_route_data, _db_run_id, _run_time)
            else:
                _bq_pusher.push_fare_snapshot(all_route_data, _db_run_id, _run_time)
                if changes:
                    _bq_pusher.push_change_events(changes, _run_time)
        except Exception as _bq_err:
            logger.warning("  [BQ] Push failed (non-fatal): %s", _bq_err)

    # Run Summary
    elapsed = _time.time() - start_time
    minutes = int(elapsed // 60)
    seconds = int(elapsed % 60)

    total_changes = sum(len(rc) for rc in changes.values()) if changes else 0
    change_counts = {}
    if changes:
        for rc in changes.values():
            for ci in rc.values():
                # Tax changes are lists of change dicts; fare changes are plain dicts
                items = ci if isinstance(ci, list) else [ci]
                for item in items:
                    if isinstance(item, dict):
                        ct = item.get("type", "unknown")
                        change_counts[ct] = change_counts.get(ct, 0) + 1

    logger.info("")
    logger.info("=" * 60)
    logger.info("  RUN SUMMARY")
    logger.info("-" * 60)
    logger.info(f"  Report:   {result_path}")
    logger.info(
        f"  Items:    {len(all_route_data)} {'airports' if args.tax else 'routes'} parsed"
    )
    if failed_commands:
        logger.info(
            f"  Failed:   {len(failed_commands)} ({', '.join(failed_commands)})"
        )
    if total_changes:
        parts = [f"{v} {k}" for k, v in sorted(change_counts.items())]
        logger.info(f"  Changes:  {total_changes} ({', '.join(parts)})")
    else:
        logger.info("  Changes:  None")
    logger.info(f"  Duration: {minutes}m {seconds}s")
    logger.info(f"  Log:      {log_file}")
    logger.info("=" * 60)

    return result_path


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        import traceback

        logger.error(
            "Run interrupted by KeyboardInterrupt. Smartpoint may not have had focus and Ctrl+C may have reached the console."
        )
        traceback.print_exc()
        sys.exit(130)
    except Exception as e:
        import traceback

        logger.error(f"Fatal error occurred:\n{traceback.format_exc()}")
        traceback.print_exc()
        sys.exit(1)
    finally:
        if getattr(sys, "frozen", False):
            input("\nPress Enter to exit...")
